"""
End-to-end checks for MFA DealPick. Run from the project folder:   python tests/test_flow.py
Uses a throwaway database; needs no network.
"""
import io
import os
import sys
import tempfile

os.environ["DATABASE_PATH"] = os.path.join(tempfile.mkdtemp(), "test.db")
os.environ["ADVOCATE_INVITE_CODE"] = "letmein-advocate"
os.environ["SECRET_KEY"] = "test"
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from docx import Document  # noqa: E402

import dealpick_app as A  # noqa: E402
import legal_content as lc  # noqa: E402

H = {"X-Requested-With": "DealPick"}
passed = 0


def check(cond, label):
    global passed
    if not cond:
        raise AssertionError("FAILED: " + label)
    passed += 1
    print("  ok  ", label)


def docx_text(resp):
    doc = Document(io.BytesIO(resp.data))
    parts = [p.text for p in doc.paragraphs]
    for t in doc.tables:
        for r in t.rows:
            parts += [c.text for c in r.cells]
    for s in doc.sections:
        parts += [p.text for p in s.header.paragraphs]
    return "\n".join(parts)


DEAL = {
    "buyer_name": "Arun Kumar", "buyer_rel": "S/o", "buyer_rel_name": "Murugan", "buyer_age": "34",
    "buyer_address": "12 Gandhi Street, Anna Nagar", "buyer_city": "Chennai", "buyer_state": "TN",
    "buyer_phone": "9840012345", "buyer_email": "arun@example.com", "buyer_id_type": "Aadhaar",
    "buyer_id_no": "1234 5678 9012", "make": "Maruti Suzuki", "model": "Swift", "variant": "VXi", "year": "2018",
    "reg_no": "tn 09 ab 1234", "reg_state": "TN", "chassis": "MA3EWDE1S00123456", "engine_no": "K12MN1234567",
    "colour": "White", "fuel": "Petrol", "odometer": "48,500", "owners": "2", "rc_holder": "Ravi Shankar",
    "hypothecated": "yes", "financier": "HDFC Bank", "insurance_policy": "Comprehensive", "insurance_valid": "2027-03-31",
    "known_defects": "Minor scratch on rear bumper", "accident_history": "none",
    "price": "5,50,000", "advance": "50000", "advance_date": "2026-10-01", "advance_mode": "UPI",
    "balance_due_date": "2026-10-20", "balance_mode": "Bank transfer", "delivery_date": "2026-10-21",
    "warranty_days": "90", "warranty_km": "3000", "transfer_cost_bearer": "Buyer", "interest_rate": "12",
    "forfeit_pct": "10", "deal_state": "TN", "place": "Chennai", "jurisdiction_city": "Chennai", "arbitration": "1",
    "agreement_date": "2026-10-01", "stamp_value": "200",
}


def run():
    c = A.app.test_client()
    adv = A.app.test_client()
    other = A.app.test_client()

    print("Accounts")
    r = c.post("/api/register", json={"role": "dealer"})
    check(r.status_code == 403, "requests without the CSRF header are blocked")
    r = c.post("/api/register", headers=H, json={
        "role": "dealer", "email": "owner@mfamotors.in", "password": "supersecret1", "name": "Mohan",
        "dealership_name": "MFA Motors", "address": "45 GST Road", "city": "Chennai", "state": "TN", "gstin": "33ABCDE1234F1Z5"})
    check(r.get_json()["success"], "dealership registers")
    r = adv.post("/api/register", headers=H, json={"role": "advocate", "email": "adv@law.in", "password": "supersecret1",
                                                   "name": "Priya Nair", "invite_code": "wrong", "bar_enrolment": "MS/1234/2015"})
    check(r.status_code == 403, "advocate sign-up rejects a wrong invite code")
    r = adv.post("/api/register", headers=H, json={"role": "advocate", "email": "adv@law.in", "password": "supersecret1",
                                                   "name": "Priya Nair", "invite_code": "letmein-advocate", "bar_enrolment": "MS/1234/2015"})
    check(r.get_json()["success"], "advocate registers with the invite code")
    other.post("/api/register", headers=H, json={
        "role": "dealer", "email": "rival@cars.in", "password": "supersecret1", "name": "Rival",
        "dealership_name": "Rival Cars", "address": "1 Main Rd", "city": "Kochi", "state": "KL"})

    print("Deals")
    r = c.post("/api/deals", headers=H, json=DEAL)
    check(r.get_json()["success"], "deal is recorded")
    did = r.get_json()["id"]
    bad = dict(DEAL, advance="900000")
    check(not c.post("/api/deals", headers=H, json=bad).get_json()["success"], "advance above price is rejected")
    d = c.get(f"/api/deals/{did}").get_json()
    check(d["deal"]["buyer_id_no"] == "XXXX XXXX 9012", "Aadhaar is stored masked")
    check(d["deal"]["reg_no"] == "TN 09 AB 1234", "registration number is normalised")
    levels = {x["level"] for x in d["compliance"]}
    check("alert" in levels, "hypothecation raises an alert")
    check(other.get(f"/api/deals/{did}").status_code == 404, "another dealership cannot open the deal")
    check(adv.get(f"/api/deals/{did}").status_code == 403, "an advocate cannot open dealer deals")

    print("Agreement")
    r = c.post(f"/api/deals/{did}/agreement", headers=H)
    aid = r.get_json()["id"]
    r = c.get(f"/api/agreements/{aid}/docx")
    t = docx_text(r)
    check(r.status_code == 200 and "Tamil" in t, "Tamil Nadu agreement mentions Tamil read-over")
    check("Tamil Nadu Motor Vehicles Taxation Act, 1974" in t, "Tamil Nadu tax Act is cited")
    check("Rs. 5,50,000/-" in t and "Five Lakh Fifty Thousand" in t, "price in Indian grouping and words")
    check("Form 35" in t, "hypothecation clause adds Form 35")
    check("Section 11 of the Arbitration" in t, "arbitration clause included when chosen")
    check("None" not in t, "no stray 'None' from empty values")
    check("{" not in t and "}" not in t, "no unfilled braces in the agreement")
    kdeal = dict(DEAL, deal_state="KL", place="Kochi", jurisdiction_city="Ernakulam", arbitration="", buyer_state="KL", reg_state="TN",
                 hypothecated="no", reg_no="KL 07 CD 4321")
    kid = c.post("/api/deals", headers=H, json=kdeal).get_json()["id"]
    kaid = c.post(f"/api/deals/{kid}/agreement", headers=H).get_json()["id"]
    kt = docx_text(c.get(f"/api/agreements/{kaid}/docx"))
    check("Malayalam" in kt and "Kerala Stamp Act, 1959" in kt, "Kerala agreement uses Malayalam and the Kerala Stamp Act")
    check("Kerala Motor Vehicles Taxation Act, 1976" in kt, "Kerala tax Act is cited")
    check("Form 28" in kt, "cross-state sale adds the Form 28 NOC")
    check("Section 11 of the Arbitration" not in kt and "exclusive jurisdiction" in kt, "courts clause used when no arbitration")
    check("explained in Tamil" not in kt and "Indian Stamp Act" not in kt, "no Tamil Nadu clauses leak into the Kerala agreement")

    print("Notice types")
    prefill = d["notice_prefill"]
    facts = {
        "payment_default": {"due_amount": prefill["due_amount"], "due_date": "2026-10-20", "demand_days": "15"},
        "cheque_dishonour": {"cheque_no": "000123", "cheque_date": "2026-10-10", "cheque_amount": "500000", "drawee_bank": "SBI Anna Nagar",
                             "dealer_bank": "ICICI Bank GST Road", "return_memo_date": "2026-10-13", "return_reason": "Funds insufficient",
                             "intimation_date": "2026-10-13"},
        "transfer_default": {"delivered_on": "2026-10-21", "demand_days": "15"},
        "cancellation_forfeiture": {"forfeit_amount": "27500", "demand_days": "7", "refund_days": "15"},
        "challan_liability": {"delivered_on": "2026-10-21", "challan_details": "TN09 e-challan 11 Nov, Rs 500", "challan_total": "500"},
        "reply_to_complaint": {"complaint_date": "2026-11-02", "complaint_summary": "AC does not cool"},
    }
    ids = {}
    for k, f in facts.items():
        r = c.post(f"/api/deals/{did}/notices", headers=H, json={"type": k, "facts": f})
        check(r.get_json()["success"], f"notice drafts: {k}")
        nid = r.get_json()["id"]
        body = c.get(f"/api/notices/{nid}").get_json()["notice"]["body"]
        check("{" not in body and "None" not in body and "LEGAL NOTICE" in body, f"{k} has no placeholders")
        ids[k] = nid
    check(not c.post(f"/api/deals/{did}/notices", headers=H, json={"type": "cheque_dishonour", "facts": {}}).get_json()["success"],
          "required notice facts are enforced")
    n138 = c.get(f"/api/notices/{ids['cheque_dishonour']}").get_json()["notice"]
    check("Section 138" in n138["body"] and "15 days" in n138["body"], "Section 138 notice carries the statutory demand")

    print("Advocate approval gate")
    nid = ids["payment_default"]
    check(adv.get(f"/api/notices/{nid}").status_code == 404, "advocate cannot see a notice still in draft")
    r = c.post(f"/api/notices/{nid}/sent", headers=H, json={"mode": "Speed post"})
    check(r.status_code == 409, "an unapproved notice cannot be marked as sent")
    dr = c.get(f"/api/notices/{nid}/docx")
    check("DO NOT SEND" in docx_text(dr) and dr.headers["Content-Disposition"].count("DRAFT") == 1, "draft download is watermarked")
    r = c.post(f"/api/notices/{nid}/decision", headers=H, json={"action": "approve"})
    check(r.status_code == 403, "a dealer cannot approve their own notice")
    check(c.post(f"/api/notices/{nid}/submit", headers=H).get_json()["success"], "dealer sends the notice to the advocate")
    q = adv.get("/api/review").get_json()
    check(any(x["id"] == nid for x in q["queue"]), "notice appears in the advocate queue")
    r = adv.post(f"/api/notices/{nid}/decision", headers=H, json={"action": "changes"})
    check(r.status_code == 400, "request-changes needs a comment")
    adv.post(f"/api/notices/{nid}/decision", headers=H, json={"action": "changes", "comment": "State the delivery date."})
    n = c.get(f"/api/notices/{nid}").get_json()["notice"]
    check(n["status"] == "changes_requested" and "delivery date" in n["comment"], "changes requested reaches the dealer")
    c.put(f"/api/notices/{nid}", headers=H, json={"title": n["title"], "body": n["body"] + "\n\nThe vehicle was delivered on 21 October 2026."})
    check(c.post(f"/api/notices/{nid}/submit", headers=H).get_json()["success"], "dealer resubmits")
    r = adv.post(f"/api/notices/{nid}/decision", headers=H, json={"action": "approve", "comment": "Fine."})
    n = r.get_json()["notice"]
    check(n["status"] == "approved" and n["final"] and n["advocate"]["enrolment"] == "MS/1234/2015", "advocate approves")
    ft = docx_text(c.get(f"/api/notices/{nid}/docx"))
    check("DO NOT SEND" not in ft and "Priya Nair" in ft and "MS/1234/2015" in ft and "Approval code" in ft,
          "approved download is clean and carries the advocate's details")
    check("DRAFT" not in c.get(f"/api/notices/{nid}/docx").headers["Content-Disposition"], "approved file is not named DRAFT")
    c.put(f"/api/notices/{nid}", headers=H, json={"title": n["title"], "body": n["body"] + "\n\nExtra line added after approval."})
    n2 = c.get(f"/api/notices/{nid}").get_json()["notice"]
    check(n2["status"] == "draft" and not n2["final"], "editing after approval withdraws the approval")
    check(c.post(f"/api/notices/{nid}/sent", headers=H, json={"mode": "Speed post"}).status_code == 409, "cannot mark as sent after approval is withdrawn")
    check("DO NOT SEND" in docx_text(c.get(f"/api/notices/{nid}/docx")), "edited notice is watermarked again")
    c.post(f"/api/notices/{nid}/submit", headers=H)
    adv.post(f"/api/notices/{nid}/decision", headers=H, json={"action": "approve"})
    r = c.post(f"/api/notices/{nid}/sent", headers=H, json={"mode": "Speed post", "date": "2026-10-14"})
    check(r.get_json()["notice"]["status"] == "sent", "approved notice can be marked as sent")
    check(c.put(f"/api/notices/{nid}", headers=H, json={"title": "x", "body": "y"}).status_code == 409, "a sent notice is locked")
    check(len(c.get(f"/api/notices/{nid}").get_json()["notice"]["events"]) >= 8, "every step is in the audit trail")
    check(other.get(f"/api/notices/{nid}").status_code == 404, "another dealership cannot open the notice")
    r = c.post(f"/api/notices/{ids['transfer_default']}/polish", headers=H)
    check(r.status_code == 400, "AI tidy-up reports clearly when it is switched off")

    regression_checks(c, adv, did)
    print(f"\nAll {passed} checks passed.")


def regression_checks(c, adv, did):
    print("Regressions")
    # Aadhaar must be masked even when the ID type was left blank
    blank = c.post("/api/deals", headers=H, json=dict(DEAL, buyer_id_type="", buyer_id_no="1234 5678 9012")).get_json()["id"]
    check(c.get(f"/api/deals/{blank}").get_json()["deal"]["buyer_id_no"] == "XXXX XXXX 9012", "a 12-digit ID is masked even if the type is blank")
    pan = c.post("/api/deals", headers=H, json=dict(DEAL, buyer_id_type="Passport", buyer_id_no="K1234567")).get_json()["id"]
    check(c.get(f"/api/deals/{pan}").get_json()["deal"]["buyer_id_no"] == "K1234567", "other IDs are kept as typed")

    # amounts
    check(lc.amount_words(5_000_000_000) == "Five Hundred Crore", "amounts of 100 crore and above do not crash")
    check(lc.inr(550000) == "Rs. 5,50,000/-", "Indian digit grouping")

    # agreement with no advance
    z = c.post("/api/deals", headers=H, json=dict(DEAL, advance="", advance_date="")).get_json()["id"]
    zaid = c.post(f"/api/deals/{z}/agreement", headers=H).get_json()["id"]
    zbody = c.get(f"/api/agreements/{zaid}").get_json()["body"]
    check("Advance of Rs. 0" not in zbody and "full Sale Price" in zbody, "no 'advance of Rs. 0' when no advance was taken")

    # a notice quotes the agreement date typed on the deal, not the day the file was generated
    c.post(f"/api/deals/{did}/agreement", headers=H)
    nid = c.post(f"/api/deals/{did}/notices", headers=H, json={"type": "transfer_default", "facts": {
        "delivered_on": "2026-10-21", "extra_facts": "Line one\n~~\n| a | b\n2. looks numbered"}}).get_json()["id"]
    body = c.get(f"/api/notices/{nid}").get_json()["notice"]["body"]
    check("agreement dated 1 October 2026" in body, "notice uses the agreement date entered on the deal")
    check(all(l.strip() != "~~" and not l.startswith("|") for l in body.splitlines()), "typed text cannot inject page breaks or tables")

    # input validation
    r = c.post(f"/api/deals/{did}/notices", headers=H, json={"type": "transfer_default", "facts": {"delivered_on": "not-a-date"}})
    check(not r.get_json()["success"], "a notice date must be a real date")
    r = c.post(f"/api/deals/{did}/notices", headers=H, json={"type": "cancellation_forfeiture", "facts": {"forfeit_amount": "900000"}})
    check(not r.get_json()["success"], "forfeiture cannot exceed the advance received")
    r = c.post(f"/api/deals/{did}/notices", headers=H, json={"type": "transfer_default", "facts": ["x"]})
    check(r.status_code == 400 and not r.get_json()["success"], "malformed notice facts give a clean 400, not a 500")
    check(not c.post("/api/deals", headers=H, json=dict(DEAL, delivery_date="31-02-2026")).get_json()["success"], "an impossible deal date is rejected")
    check(not c.post("/api/deals", headers=H, json=dict(DEAL, year="1800")).get_json()["success"], "an impossible year is rejected")

    # non-ASCII invite code is a normal 403, not a crash
    r = adv.post("/api/register", headers=H, json={"role": "advocate", "email": "x@y.in", "password": "supersecret1", "name": "X",
                                                   "invite_code": "\u0b95\u0bcb\u0b9f\u0bcd", "bar_enrolment": "1"})
    check(r.status_code == 403, "a non-ASCII invite code is refused cleanly")

    # editing a deal's state keeps its reference in step
    e = c.post("/api/deals", headers=H, json=DEAL).get_json()["id"]
    c.put(f"/api/deals/{e}", headers=H, json=dict(DEAL, deal_state="KL"))
    check(c.get(f"/api/deals/{e}").get_json()["deal"]["ref"].startswith("DP-KL-"), "deal reference follows the state")

    # mark-as-sent date rules
    sid = c.post(f"/api/deals/{did}/notices", headers=H, json={"type": "transfer_default", "facts": {"delivered_on": "2026-10-21"}}).get_json()["id"]
    c.post(f"/api/notices/{sid}/submit", headers=H)
    adv.post(f"/api/notices/{sid}/decision", headers=H, json={"action": "approve"})
    tomorrow = (lc.today_ist() + __import__("datetime").timedelta(days=1)).isoformat()
    check(c.post(f"/api/notices/{sid}/sent", headers=H, json={"mode": "Speed post", "date": tomorrow}).status_code == 400, "a sent date in the future is refused")

    # multi-line addresses stay in one paragraph in the documents
    ml = c.post("/api/deals", headers=H, json=dict(DEAL, buyer_address="12 Gandhi Street,\nAnna Nagar", known_defects="Scratch\nDent")).get_json()["id"]
    maid = c.post(f"/api/deals/{ml}/agreement", headers=H).get_json()["id"]
    mb = c.get(f"/api/agreements/{maid}").get_json()["body"]
    check("12 Gandhi Street, Anna Nagar" in mb and "Scratch; Dent" in mb, "multi-line addresses and defects stay on one line")


if __name__ == "__main__":
    run()
