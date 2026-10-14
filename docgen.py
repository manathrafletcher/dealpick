"""
MFA DealPick - Word document builder.
Lays out the lightweight markup produced by legal_content.py (see that file's docstring).
Documents are built in memory and streamed to the browser; nothing is written to disk.
"""

import io
import re

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

FONT = "Times New Roman"
_NUM = re.compile(r"^(\d{1,3})[\.\)]\s+(.*)$")
_SUB = re.compile(r"^\(([a-zA-Z]{1,3}|\d{1,2})\)\s+(.*)$")
_INLINE = re.compile(r"(\*\*[^*\n]+?\*\*)")


def _style(run, size=12, bold=False, italic=False, color=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold or None
    run.italic = italic or None
    if color:
        run.font.color.rgb = RGBColor(*color)


def _runs(p, text, size=12, bold=False):
    for part in _INLINE.split(text):
        if not part:
            continue
        b = bold
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            part, b = part[2:-2], True
        _style(p.add_run(part.replace("**", "")), size=size, bold=b)


def _fmt(p, before=4, after=4, left=0.0, first=0.0, align=WD_ALIGN_PARAGRAPH.JUSTIFY, keep_next=False):
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.left_indent, pf.first_line_indent = Inches(left), Inches(first)
    pf.keep_with_next = keep_next
    p.alignment = align


def _field(p, instr, size=9):
    for kind in ("begin", None, "separate", "text", "end"):
        r = p.add_run()
        _style(r, size=size)
        if kind in ("begin", "separate", "end"):
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
            r._r.append(el)
        elif kind is None:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = instr
            r._r.append(el)
        else:
            r.text = "1"


def _table(doc, rows, bordered):
    cols = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=cols)
    if bordered:
        t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, r in enumerate(rows):
        for j in range(cols):
            cell = t.cell(i, j)
            para = cell.paragraphs[0]
            _fmt(para, 3, 3, align=WD_ALIGN_PARAGRAPH.LEFT)
            _runs(para, r[j] if j < len(r) else "", size=11 if bordered else 12)
    if cols == 2 and bordered:
        for row in t.rows:
            row.cells[0].width, row.cells[1].width = Inches(2.1), Inches(3.9)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build_docx(text, header_banner=None, banner_color=(0xB4, 0x32, 0x5A), footer="MFA DealPick"):
    doc = Document()
    for sec in doc.sections:
        sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)      # A4
        sec.top_margin = sec.bottom_margin = Inches(1)
        sec.left_margin = sec.right_margin = Inches(1.1)
        if header_banner:
            hp = sec.header.paragraphs[0]
            _fmt(hp, 0, 0, align=WD_ALIGN_PARAGRAPH.CENTER)
            _style(hp.add_run(header_banner), size=10, bold=True, color=banner_color)
        fp = sec.footer.paragraphs[0]
        _fmt(fp, 0, 0, align=WD_ALIGN_PARAGRAPH.CENTER)
        _style(fp.add_run(f"{footer}   Page "), size=9)
        _field(fp, "PAGE")
        _style(fp.add_run(" of "), size=9)
        _field(fp, "NUMPAGES")

    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = FONT, Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    lines = text.split("\n")
    i = 0
    while i < len(lines):
        raw = lines[i].rstrip()
        s = raw.strip()
        i += 1
        if s == "~~":
            doc.add_page_break()
            continue
        if s.startswith("|"):                                  # table block
            rows, bordered = [], not s.startswith("|~")
            i -= 1
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = lines[i].strip()
                cells = cells[2:] if cells.startswith("|~") else cells[1:]
                rows.append([c.strip().lstrip("~") for c in cells.split("|")])
                i += 1
            _table(doc, rows, bordered)
            continue
        if not s:
            continue
        p = doc.add_paragraph()
        if s.startswith("### "):
            _fmt(p, 10, 3, align=WD_ALIGN_PARAGRAPH.LEFT, keep_next=True)
            _runs(p, s[4:], bold=True)
        elif s.startswith("## "):
            _fmt(p, 8, 6, align=WD_ALIGN_PARAGRAPH.CENTER, keep_next=True)
            _runs(p, s[3:], size=13, bold=True)
        elif s.startswith("# "):
            _fmt(p, 0, 14, align=WD_ALIGN_PARAGRAPH.CENTER, keep_next=True)
            _runs(p, s[2:], size=16, bold=True)
        elif s.startswith(">>"):
            _fmt(p, 0, 2, align=WD_ALIGN_PARAGRAPH.RIGHT)
            _runs(p, s[2:].strip())
        elif _NUM.match(s):
            m = _NUM.match(s)
            _fmt(p, 4, 4, left=0.4, first=-0.4)
            _style(p.add_run(f"{m.group(1)}.  "), bold=False)
            _runs(p, m.group(2))
        elif _SUB.match(s):
            m = _SUB.match(s)
            _fmt(p, 3, 3, left=0.8, first=-0.4)
            _style(p.add_run(f"({m.group(1)})  "))
            _runs(p, m.group(2))
        elif s.startswith("- "):
            _fmt(p, 2, 2, left=0.7, first=-0.25)
            _style(p.add_run("\u2022  "))
            _runs(p, s[2:])
        else:
            _fmt(p, 5, 5)
            _runs(p, s)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf
