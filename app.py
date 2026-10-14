"""
MFA DealPick - one-time-use prototype (single file)
---------------------------------------------------
A used-car dealership logs in, enters a deal (Tamil Nadu or Kerala), gets an AI review of the deal,
answers a few clarifying questions, and downloads the legal drafts for that deal in ONE Word file.

Nothing is stored: no database, no files. Each request carries everything it needs, and the
Word file is built in memory and streamed to the browser.

Sections of this file
  1. AI client (Groq free tier)       4. Deal / notice validation
  2. Word document builder            5. AI deal review and clarifying questions
  3. Legal content (TN/KL templates)  6. Web routes

Files:  app.py (this file) + index.html (the whole front end).

Run locally:  python app.py                      (http://127.0.0.1:8081)
Deploy:       gunicorn app:app --workers 1 --threads 8 --timeout 180

Environment variables
  DEALER_USERNAME / DEALER_PASSWORD   prototype login (defaults: demo / demo1234; change them)
  GROQ_API_KEY                        switches on the AI review and AI questions (optional)
  SECRET_KEY                          signs the login cookie (optional; random per start if unset)

These are drafting aids, not legal advice. Every document must be reviewed by an advocate before use.
"""

import datetime as dt
import hmac
import io
import json
import os
import re
import secrets
import threading
import time
import traceback

import requests
from requests.adapters import HTTPAdapter
from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from flask import Flask, Response, g, jsonify, request, send_file, session
from werkzeug.exceptions import HTTPException
from werkzeug.middleware.proxy_fix import ProxyFix

# ============================================================================
# 1. AI CLIENT (Groq, free tier): model discovery, failover, small token budgets
# ============================================================================

GROQ_BASE_URL = os.environ.get("GROQ_BASE_URL", "https://api.groq.com/openai/v1").rstrip("/")

_PREFERRED = ["llama-3.3-70b-versatile", "openai/gpt-oss-120b",
              "llama-3.1-8b-instant", "openai/gpt-oss-20b"]
_EXCLUDED = ("whisper", "guard", "safeguard", "compound", "tts", "orpheus",
             "playai", "embed", "moderation")
MAX_MODELS_TRIED = 4
RATE_LIMIT_MAX_WAIT = 12

_HTTP = requests.Session()
_HTTP.mount("https://", HTTPAdapter(pool_connections=4, pool_maxsize=8))

_CACHE = {"models": None, "expires": 0.0}
_LOCK = threading.Lock()


class _Skip(Exception):
    pass


class _RateLimited(Exception):
    def __init__(self, wait):
        super().__init__(f"rate limited, retry in {wait}s")
        self.wait = wait


def ai_available() -> bool:
    return bool(os.environ.get("GROQ_API_KEY", "").strip())


def _models(api_key):
    now = time.time()
    with _LOCK:
        if _CACHE["models"] is None or now > _CACHE["expires"]:
            ids, ok = [], False
            try:
                r = _HTTP.get(f"{GROQ_BASE_URL}/models",
                              headers={"Authorization": f"Bearer {api_key}"}, timeout=(4, 6))
                if r.status_code == 200:
                    ids = [i["id"] for i in r.json().get("data", [])
                           if isinstance(i, dict) and i.get("id") and i.get("active", True)]
                    ok = True
            except Exception as e:                      # discovery is best-effort
                print(f"[AI] model discovery failed: {e}")
            if ids:
                sel = [m for m in _PREFERRED if m in ids]
                sel += [m for m in ids if m not in sel and not any(x in m.lower() for x in _EXCLUDED)]
            else:
                sel = list(_PREFERRED)
            override = os.environ.get("GROQ_MODEL", "").strip()
            if override:
                sel = [override] + [m for m in sel if m != override]
            _CACHE["models"] = list(dict.fromkeys(sel))
            _CACHE["expires"] = now + (3600 if ok else 30)
        return list(_CACHE["models"])[:MAX_MODELS_TRIED]


def _once(api_key, model, messages, temperature, max_tokens):
    payload = {"model": model, "messages": messages, "temperature": temperature,
               "max_completion_tokens": max_tokens}
    if model.startswith("openai/gpt-oss"):
        payload["reasoning_effort"] = "low"
    try:
        r = _HTTP.post(f"{GROQ_BASE_URL}/chat/completions",
                       headers={"Authorization": f"Bearer {api_key}",
                                "Content-Type": "application/json"},
                       json=payload, timeout=(5, 60))
    except requests.exceptions.RequestException as e:
        raise _Skip(f"request error on {model}: {e}")
    if r.status_code == 429:
        try:
            wait = max(0.5, min(float(r.headers.get("retry-after", "")), 60.0))
        except (TypeError, ValueError):
            wait = 5.0
        raise _RateLimited(wait)
    if r.status_code != 200:
        raise _Skip(f"HTTP {r.status_code} on {model}: {r.text[:160]}")
    try:
        choice = r.json()["choices"][0]
        text = choice["message"]["content"] or ""
    except Exception as e:
        raise _Skip(f"bad response from {model}: {e}")
    if not text.strip():
        raise _Skip(f"empty content from {model}")
    return text, choice.get("finish_reason")


def ai_generate(prompt: str, system: str = "", temperature: float = 0.2,
                max_tokens: int = 2800) -> str:
    api_key = os.environ.get("GROQ_API_KEY", "").strip()
    if not api_key:
        raise RuntimeError("AI is not set up. Add GROQ_API_KEY in the server settings.")
    messages = ([{"role": "system", "content": system}] if system else []) + \
               [{"role": "user", "content": prompt}]
    models = _models(api_key)
    last = None
    for rnd in range(2):
        wait = None
        for m in models:
            try:
                text, finish = _once(api_key, m, messages, temperature, max_tokens)
                if finish == "length":
                    raise _Skip(f"{m} ran out of tokens before finishing")
                return text.strip()
            except _RateLimited as e:
                last = f"rate limited on {m}"
                wait = e.wait if wait is None else min(wait, e.wait)
            except _Skip as e:
                last = str(e)
        if wait is None or rnd == 1:
            break
        time.sleep(min(wait, RATE_LIMIT_MAX_WAIT))
    raise RuntimeError(f"The AI service could not finish this request ({last}).")


# ============================================================================
# 2. WORD DOCUMENT BUILDER (lays out the markup described in section 3)
# ============================================================================

FONT = "Times New Roman"
_NUM = re.compile(r"^(\d{1,3})[\.\)]\s+(.*)$")
_SUB = re.compile(r"^\(([a-zA-Z]{1,3}|\d{1,2})\)\s+(.*)$")
_INLINE = re.compile(r"(\*\*[^*\n]+?\*\*)")


def _style(run, size=12, bold=False, italic=False, color=None):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = bold or None
    run.italic = italic or None
    if color:
        run.font.color.rgb = RGBColor(*color)


def _runs(p, text, size=12, bold=False):
    for part in _INLINE.split(text):
        if not part:
            continue
        b = bold
        if part.startswith("**") and part.endswith("**") and len(part) > 4:
            part, b = part[2:-2], True
        _style(p.add_run(part.replace("**", "")), size=size, bold=b)


def _fmt(p, before=4, after=4, left=0.0, first=0.0, align=WD_ALIGN_PARAGRAPH.JUSTIFY, keep_next=False):
    pf = p.paragraph_format
    pf.space_before, pf.space_after = Pt(before), Pt(after)
    pf.left_indent, pf.first_line_indent = Inches(left), Inches(first)
    pf.keep_with_next = keep_next
    p.alignment = align


def _docx_field(p, instr, size=9):
    for kind in ("begin", None, "separate", "text", "end"):
        r = p.add_run()
        _style(r, size=size)
        if kind in ("begin", "separate", "end"):
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
            r._r.append(el)
        elif kind is None:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = instr
            r._r.append(el)
        else:
            r.text = "1"


def _table(doc, rows, bordered):
    cols = max(len(r) for r in rows)
    t = doc.add_table(rows=len(rows), cols=cols)
    if bordered:
        t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, r in enumerate(rows):
        for j in range(cols):
            cell = t.cell(i, j)
            para = cell.paragraphs[0]
            _fmt(para, 3, 3, align=WD_ALIGN_PARAGRAPH.LEFT)
            _runs(para, r[j] if j < len(r) else "", size=11 if bordered else 12)
    if cols == 2 and bordered:
        for row in t.rows:
            row.cells[0].width, row.cells[1].width = Inches(2.1), Inches(3.9)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def build_docx(text, header_banner=None, banner_color=(0xB4, 0x32, 0x5A), footer="MFA DealPick"):
    doc = Document()
    for sec in doc.sections:
        sec.page_width, sec.page_height = Inches(8.27), Inches(11.69)      # A4
        sec.top_margin = sec.bottom_margin = Inches(1)
        sec.left_margin = sec.right_margin = Inches(1.1)
        if header_banner:
            hp = sec.header.paragraphs[0]
            _fmt(hp, 0, 0, align=WD_ALIGN_PARAGRAPH.CENTER)
            _style(hp.add_run(header_banner), size=10, bold=True, color=banner_color)
        fp = sec.footer.paragraphs[0]
        _fmt(fp, 0, 0, align=WD_ALIGN_PARAGRAPH.CENTER)
        _style(fp.add_run(f"{footer}   Page "), size=9)
        _docx_field(fp, "PAGE")
        _style(fp.add_run(" of "), size=9)
        _docx_field(fp, "NUMPAGES")

    normal = doc.styles["Normal"]
    normal.font.name, normal.font.size = FONT, Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    lines = text.split("\n")
    i = 0
    while i < len(lines):
        raw = lines[i].rstrip()
        s = raw.strip()
        i += 1
        if s == "~~":
            doc.add_page_break()
            continue
        if s.startswith("|"):                                  # table block
            rows, bordered = [], not s.startswith("|~")
            i -= 1
            while i < len(lines) and lines[i].strip().startswith("|"):
                cells = lines[i].strip()
                cells = cells[2:] if cells.startswith("|~") else cells[1:]
                rows.append([c.strip().lstrip("~") for c in cells.split("|")])
                i += 1
            _table(doc, rows, bordered)
            continue
        if not s:
            continue
        p = doc.add_paragraph()
        if s.startswith("### "):
            _fmt(p, 10, 3, align=WD_ALIGN_PARAGRAPH.LEFT, keep_next=True)
            _runs(p, s[4:], bold=True)
        elif s.startswith("## "):
            _fmt(p, 8, 6, align=WD_ALIGN_PARAGRAPH.CENTER, keep_next=True)
            _runs(p, s[3:], size=13, bold=True)
        elif s.startswith("# "):
            _fmt(p, 0, 14, align=WD_ALIGN_PARAGRAPH.CENTER, keep_next=True)
            _runs(p, s[2:], size=16, bold=True)
        elif s.startswith(">>"):
            _fmt(p, 0, 2, align=WD_ALIGN_PARAGRAPH.RIGHT)
            _runs(p, s[2:].strip())
        elif _NUM.match(s):
            m = _NUM.match(s)
            _fmt(p, 4, 4, left=0.4, first=-0.4)
            _style(p.add_run(f"{m.group(1)}.  "), bold=False)
            _runs(p, m.group(2))
        elif _SUB.match(s):
            m = _SUB.match(s)
            _fmt(p, 3, 3, left=0.8, first=-0.4)
            _style(p.add_run(f"({m.group(1)})  "))
            _runs(p, m.group(2))
        elif s.startswith("- "):
            _fmt(p, 2, 2, left=0.7, first=-0.25)
            _style(p.add_run("\u2022  "))
            _runs(p, s[2:])
        else:
            _fmt(p, 5, 5)
            _runs(p, s)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf


# ============================================================================
# 3. LEGAL CONTENT: state-aware templates for TN and KL (sale agreement + six notice types)
#    Markup: # Title | ## Heading | ### Clause head | N. numbered | (a) sub-item | - bullet
#            |a|b bordered row | |~a|b borderless row | >> right-aligned | ~~ page break | **bold**
# ============================================================================

import datetime as _dt

_IST = _dt.timezone(_dt.timedelta(hours=5, minutes=30))

STATES = {
    "TN": {
        "name": "Tamil Nadu",
        "authority": "Regional Transport Office (RTO) of the Transport Department, Government of Tamil Nadu",
        "authority_short": "jurisdictional Regional Transport Office (RTO) in Tamil Nadu",
        "tax_act": "the Tamil Nadu Motor Vehicles Taxation Act, 1974",
        "stamp_act": "the Indian Stamp Act, 1899 as in force in the State of Tamil Nadu",
        "hc": "the High Court of Judicature at Madras",
        "lang": "Tamil",
        "default_city": "Chennai",
    },
    "KL": {
        "name": "Kerala",
        "authority": "Regional Transport Office (RTO) or Sub Regional Transport Office of the Motor Vehicles Department, Government of Kerala",
        "authority_short": "jurisdictional Regional Transport Office (RTO) or Sub RTO in Kerala",
        "tax_act": "the Kerala Motor Vehicles Taxation Act, 1976",
        "stamp_act": "the Kerala Stamp Act, 1959",
        "hc": "the High Court of Kerala",
        "lang": "Malayalam",
        "default_city": "Kochi",
    },
}
STATE_LABELS = {"TN": "Tamil Nadu", "KL": "Kerala", "OT": "another state"}

# ---------------------------------------------------------------- formatting helpers

_ONES = ["", "One", "Two", "Three", "Four", "Five", "Six", "Seven", "Eight", "Nine", "Ten",
         "Eleven", "Twelve", "Thirteen", "Fourteen", "Fifteen", "Sixteen", "Seventeen",
         "Eighteen", "Nineteen"]
_TENS = ["", "", "Twenty", "Thirty", "Forty", "Fifty", "Sixty", "Seventy", "Eighty", "Ninety"]


def _two(n):
    return _ONES[n] if n < 20 else (_TENS[n // 10] + (" " + _ONES[n % 10] if n % 10 else ""))


def amount_words(n):
    n = int(round(float(n or 0)))
    if n == 0:
        return "Zero"
    parts = []
    for div, name in ((10 ** 7, "Crore"), (10 ** 5, "Lakh"), (1000, "Thousand"), (100, "Hundred")):
        q, n = divmod(n, div)
        if q:
            parts.append(f"{amount_words(q) if q >= 100 else _two(q)} {name}")
    if n:
        parts.append(_two(n))
    return " ".join(parts)


def num(v, default=0.0):
    try:
        return float(str(v).replace(",", "").strip())
    except (TypeError, ValueError):
        return default


def inr(n):
    n = int(round(num(n)))
    s = str(abs(n))
    if len(s) > 3:
        head, tail, groups = s[:-3], s[-3:], []
        while len(head) > 2:
            groups.insert(0, head[-2:])
            head = head[:-2]
        if head:
            groups.insert(0, head)
        s = ",".join(groups + [tail])
    return ("-" if n < 0 else "") + "Rs. " + s + "/-"


def inr_words(n):
    return f"{inr(n)} (Rupees {amount_words(n)} only)"


def today_ist():
    """Today's date in India. Servers run on UTC, which is a day behind India for the first 5.5 hours of each day."""
    return _dt.datetime.now(_IST).date()


def ist_date(ts):
    """'2026-10-13T20:00:00Z' (UTC timestamp) -> '2026-10-14' (the date in India)."""
    try:
        t = _dt.datetime.strptime(str(ts)[:19], "%Y-%m-%dT%H:%M:%S").replace(tzinfo=_dt.timezone.utc)
        return t.astimezone(_IST).date().isoformat()
    except (TypeError, ValueError):
        return ""


def dmy_utc(ts):
    return dmy(ist_date(ts))


def dmy(iso):
    try:
        return _dt.date.fromisoformat(str(iso)[:10]).strftime("%d %B %Y").lstrip("0")
    except (TypeError, ValueError):
        return ""


def parse_date(iso):
    try:
        return _dt.date.fromisoformat(str(iso)[:10])
    except (TypeError, ValueError):
        return None


def _one_line(v, sep=" "):
    """Collapse line breaks and runs of spaces, so typed text cannot start a new paragraph, table row or page break."""
    parts = [" ".join(x.split()) for x in str(v or "").splitlines()]
    if sep.startswith(","):
        parts = [x.rstrip(",") for x in parts]
    return sep.join(x for x in parts if x)


_MULTILINE = {"extra_clauses"}                       # one term per line: these keep their line breaks
_ADDRESS = {"buyer_address", "address"}


def _s(d, k, default=""):
    v = str(d.get(k) or default).strip()
    if k in _MULTILINE:
        return v
    return _one_line(v, ", " if k in _ADDRESS else ("; " if k == "known_defects" else " "))


def _cell(v):
    """Table cells are split on '|', so a typed '|' would add a column."""
    return str(v).replace("|", "/")


def vehicle_label(d):
    return " ".join(x for x in (_s(d, "year"), _s(d, "make"), _s(d, "model"), _s(d, "variant")) if x)


def vehicle_long(d):
    bits = [vehicle_label(d), f"registration no. {_s(d, 'reg_no')}"]
    if _s(d, "chassis"):
        bits.append(f"chassis no. {_s(d, 'chassis')}")
    if _s(d, "engine_no"):
        bits.append(f"engine no. {_s(d, 'engine_no')}")
    return ", ".join(b for b in bits if b.strip())


def deal_state(profile, d):
    st = _s(d, "deal_state") or profile.get("state") or "TN"
    return st if st in STATES else "TN"


def balance_due(d):
    return max(0.0, num(d.get("price")) - num(d.get("advance")))


def buyer_block(d, short=False):
    rel = f"{_s(d, 'buyer_rel', 'S/o')} {_s(d, 'buyer_rel_name')}".strip()
    age = f", aged {_s(d, 'buyer_age')} years" if _s(d, "buyer_age") else ""
    addr = ", ".join(x for x in (_s(d, "buyer_address"), _s(d, "buyer_city")) if x)
    if short:
        return f"{_s(d, 'buyer_name')}, {rel}{age}, residing at {addr}"
    idn = f", holding {_s(d, 'buyer_id_type')} {_s(d, 'buyer_id_no')}" if _s(d, "buyer_id_no") else ""
    ph = f", mobile {_s(d, 'buyer_phone')}" if _s(d, "buyer_phone") else ""
    return f"{_s(d, 'buyer_name')}, {rel}{age}, residing at {addr}{idn}{ph}"


def seller_block(p, st):
    S = STATES[st]
    bits = [f"**{p['dealership_name']}**, a dealer in pre-owned motor vehicles with its place of business at "
            f"{_s(p, 'address')}, {p['city']}, {S['name']}"]
    if p.get("gstin"):
        bits.append(f"GSTIN {p['gstin']}")
    if p.get("trade_cert"):
        bits.append(f"trade certificate no. {p['trade_cert']}")
    who = p.get("signatory") or "its authorised signatory"
    return ", ".join(bits) + f", represented by {who}"


# ---------------------------------------------------------------- compliance checks

def compliance(p, d, today=None):
    """Return [{level, text}] where level is info | warn | alert."""
    today = today or today_ist()
    st = deal_state(p, d)
    S = STATES[st]
    out = []

    def add(level, text):
        out.append({"level": level, "text": text})

    add("info", f"Hand over Form 29 (signed by the registered owner) and Form 30 (signed by the buyer), "
                f"two sets each. The buyer must report the transfer to the {S['authority_short']} within "
                f"30 days; you report it within 14 days (Section 50, Motor Vehicles Act, 1988).")
    add("info", f"Print the agreement on stamp paper or an e-stamp bought in {S['name']}. Your advocate "
                f"confirms the value under {S['stamp_act']}; type it into the deal so it appears in the agreement.")
    add("info", f"Read the agreement out to the buyer in {S['lang']} before signing. The agreement records that you did.")

    reg, bst = _s(d, "reg_state"), _s(d, "buyer_state")
    if (reg and bst and reg != bst) or (reg and reg not in ("TN", "KL") and reg != st):
        add("warn", "The vehicle is registered in a different state from the buyer's home. You need an NOC "
                    "(Form 28) from the registering authority, and the buyer must re-register the vehicle in "
                    "their own state within 12 months of moving it (Section 47).")
    if _s(d, "hypothecated") == "yes":
        fin = _s(d, "financier") or "the financier"
        add("alert", f"The RC shows a loan with {fin}. Close the loan and get Form 35 and the financier's NOC "
                     f"before handover. The agreement makes this your responsibility.")
    if _s(d, "finance_by_buyer") == "yes":
        add("info", "The buyer is financing the purchase. Confirm the lender's disbursal date matches the balance "
                    "due date, and cooperate on the hypothecation entry (Form 34).")
    yr = int(num(d.get("year"))) if num(d.get("year")) else 0
    if yr and today.year - yr >= 13:
        add("warn", "This car is 13 or more years old. Registration lasts 15 years from first registration, then "
                    "needs renewal (Section 41(7)). Check the validity date on the RC before you promise delivery.")
    cash_hit = ((_s(d, "advance_mode") == "Cash" and num(d.get("advance")) >= 200000) or
                (_s(d, "balance_mode") == "Cash" and balance_due(d) >= 200000))
    if cash_hit:
        add("alert", "Income-tax law restricts taking Rs. 2,00,000 or more in cash. Take a bank transfer, "
                     "cheque or demand draft instead.")
    if num(d.get("price")) > 1000000:
        add("info", "A sale above Rs. 10,00,000 may attract tax collected at source. Confirm with your accountant "
                    "before you issue the invoice.")
    ins, dlv = parse_date(d.get("insurance_valid")), parse_date(d.get("delivery_date"))
    if ins and dlv and ins < dlv:
        add("warn", "The insurance policy expires before the delivery date. Renew it or tell the buyer in writing.")
    age = int(num(d.get("buyer_age"))) if num(d.get("buyer_age")) else 0
    if 0 < age < 18:
        add("alert", "The buyer is under 18. A minor cannot make a binding contract. Put the parent or guardian in as the buyer "
                     "and have them sign.")
    if _s(d, "accident_history") == "major":
        add("info", "A major repair is on record. Describe it under known defects so the agreement shows it was disclosed.")
    return out


# ---------------------------------------------------------------- sale agreement

def build_agreement(p, d, today=None):
    today = today or today_ist()
    st = deal_state(p, d)
    S = STATES[st]
    price, adv, bal = num(d.get("price")), num(d.get("advance")), balance_due(d)
    place = _s(d, "place") or p.get("city") or S["default_city"]
    jur = _s(d, "jurisdiction_city") or place
    agr_date = dmy(d.get("agreement_date")) or dmy(today.isoformat())
    rate = _s(d, "interest_rate", "12")
    pct = _s(d, "forfeit_pct", "10")
    reg_st = _s(d, "reg_state")
    cross = bool(reg_st and _s(d, "buyer_state") and reg_st != _s(d, "buyer_state"))
    hypo = _s(d, "hypothecated") == "yes"
    rc_holder = _s(d, "rc_holder") or p["dealership_name"]

    L = []
    n = [0]

    def clause(title):
        n[0] += 1
        L.append(f"### {n[0]}. {title}")

    def sub(items):
        for i, t in enumerate(items):
            L.append(f"({chr(97 + i)}) {t}")

    L += ["# AGREEMENT FOR SALE OF A PRE-OWNED MOTOR VEHICLE", ""]
    L.append(f"This Agreement is made at {place}, {S['name']}, on {agr_date}, between:")
    L.append(f"{seller_block(p, st)} (the \"**Seller**\", which includes its successors and assigns); and")
    L.append(f"{buyer_block(d)} (the \"**Buyer**\", which includes the Buyer's heirs, legal representatives and assigns).")
    L.append("**Whereas:**")
    L.append(f"(A) The Seller deals in pre-owned motor vehicles and is in lawful possession of the vehicle described in "
             f"Clause 1 (the \"**Vehicle**\"), whose registration certificate stands in the name of {rc_holder}.")
    L.append("(B) The Buyer has inspected the Vehicle, taken a test drive, and wishes to buy it. The Seller wishes to "
             "sell it. Both agree on the following terms.")
    L.append("**The parties agree as follows:**")

    # 1 vehicle
    clause("The Vehicle")
    rows = [("Make and model", " ".join(x for x in (_s(d, "make"), _s(d, "model"), _s(d, "variant")) if x)),
            ("Year of manufacture", _s(d, "year")),
            ("Registration no.", _s(d, "reg_no")),
            ("Chassis no.", _s(d, "chassis")),
            ("Engine no.", _s(d, "engine_no")),
            ("Colour and fuel", ", ".join(x for x in (_s(d, "colour"), _s(d, "fuel")) if x)),
            ("Odometer reading at sale", f"{_s(d, 'odometer')} km" if _s(d, "odometer") else ""),
            ("Owner serial", _s(d, "owners")),
            ("RC in the name of", rc_holder),
            ("Insurance", ", ".join(x for x in (_s(d, "insurance_policy"),
                                               f"valid to {dmy(d.get('insurance_valid'))}" if dmy(d.get("insurance_valid")) else "") if x)),
            ("Hypothecation", f"Subsisting loan with {_s(d, 'financier') or 'a financier'}" if hypo else "No loan or hypothecation")]
    for k, v in rows:
        if v:
            L.append(f"|**{k}**|{_cell(v)}")
    L.append("")

    # 2 price
    clause("Sale price and payment")
    L.append(f"The Seller sells and the Buyer buys the Vehicle for **{inr_words(price)}** (the \"Sale Price\"), payable as follows:")
    if adv > 0:
        sub([f"Advance of {inr_words(adv)} paid on {dmy(d.get('advance_date')) or 'the date of this Agreement'} by "
             f"{_s(d, 'advance_mode', 'bank transfer')}, receipt of which the Seller acknowledges.",
             f"Balance of {inr_words(bal)} payable on or before {dmy(d.get('balance_due_date')) or 'the delivery date'} by "
             f"{_s(d, 'balance_mode', 'bank transfer')}, against delivery of the Vehicle."])
    else:                                              # no advance was taken: the whole price is the balance
        sub([f"The full Sale Price of {inr_words(price)} is payable on or before {dmy(d.get('balance_due_date')) or 'the delivery date'} by "
             f"{_s(d, 'balance_mode', 'bank transfer')}, against delivery of the Vehicle."])
    if _s(d, "finance_by_buyer") == "yes":
        L.append(f"The Buyer is financing the Balance through {_s(d, 'buyer_financier') or 'a financier'}. The Buyer shall "
                 f"ensure the financier pays the Seller on or before the due date. A delay by the financier is a delay by the Buyer.")
    L.append("The Sale Price includes GST, if any, payable by the Seller on its margin under the Central Goods and Services "
             "Tax Act, 2017. The Seller shall issue a bill of sale or tax invoice. Any tax that must be collected at source "
             "under the income-tax law shall be collected and accounted for as that law requires. Payments of Rs. 2,00,000 "
             "or more shall be made through banking channels and not in cash.")

    # 3 delivery
    clause("Delivery, documents, risk and title")
    L.append(f"The Seller shall deliver the Vehicle with its keys at {place} on {dmy(d.get('delivery_date')) or 'the date on which the Balance is received'}, "
             f"on receipt of the full Sale Price, together with:")
    docs = ["the original registration certificate" + (" (if the financier still holds it, a copy with the financier's written undertaking to release the original)" if hypo else ""),
            "Form 29 signed by the registered owner, two sets, and Form 30 signed by the Buyer",
            "the valid insurance policy and pollution-under-control certificate",
            f"proof of payment of tax under {S['tax_act']}",
            "the service history and all keys in the Seller's possession"]
    if hypo:
        docs.append("Form 35 and the financier's no-objection certificate")
    if cross:
        docs.append("a no-objection certificate in Form 28 from the registering authority")
    sub(docs)
    L.append("Risk in the Vehicle passes to the Buyer on delivery. Property in the Vehicle passes to the Buyer on delivery "
             "against full payment of the Sale Price.")

    # 4 condition
    clause("Condition of the Vehicle")
    L.append("The Buyer has inspected the Vehicle and its documents and accepts its mechanical, electrical and cosmetic "
             "condition as it stands (\"as is where is\"), subject to Clauses 5 and 12.")
    L.append(f"The Seller has disclosed the following known defects: {_s(d, 'known_defects') or 'none, other than normal wear and tear for its age and use'}.")
    acc = _s(d, "accident_history")
    L.append({"major": "The Seller has also disclosed that the Vehicle has had a major repair after an accident, as described above.",
              "minor": "The Seller has disclosed that the Vehicle has had minor cosmetic or bodywork repairs.",
              }.get(acc, "To the Seller's knowledge, the Vehicle has not been in a major accident."))

    # 5 seller assurances
    clause("Seller's assurances")
    L.append("The Seller assures the Buyer that:")
    sub(["it has the right to sell the Vehicle, or the written authority of the registered owner to do so;",
         "the Vehicle is free from any lien, attachment, hypothecation or other encumbrance, other than any loan "
         "disclosed in Clause 1 that the Seller will close under Clause 9;",
         "the Vehicle is not stolen and is not under seizure by any police authority or court;",
         f"all tax due under {S['tax_act']}, and all fines and challans (including electronic challans) relating to the period "
         f"up to delivery, are paid or will be paid by the Seller;",
         "to the Seller's knowledge, the odometer reading is genuine and the Vehicle has not been submerged in a flood "
         "or written off as a total loss, except as disclosed."])
    L.append("If an assurance in (a), (b) or (c) proves untrue, the Seller shall remove the defect at its cost within 15 days of "
             "written notice. If it cannot, the Seller shall take the Vehicle back and refund everything it has received.")

    # 6 buyer assurances
    clause("Buyer's assurances")
    L.append("The Buyer assures the Seller that the Buyer holds a valid driving licence, has paid the Sale Price from lawful funds, "
             "will use the Vehicle lawfully, and will tell the Seller in writing within 7 days if the Buyer's address changes.")

    # 7 transfer
    clause("Transfer of ownership")
    items = [f"The Seller shall report the transfer to the {S['authority']}, within 14 days of delivery, and send the Buyer a copy of the report (Section 50, Motor Vehicles Act, 1988).",
             f"The Buyer shall apply to the {S['authority_short']} within 30 days of delivery, with Form 29, Form 30, the registration certificate, identity and address proof, insurance, and the fees. The Buyer shall give the Seller the acknowledgment within 7 days of applying, and a copy of the updated registration certificate when issued."]
    if cross:
        items.append(f"The Vehicle is registered in {STATE_LABELS.get(reg_st, reg_st)} and the Buyer lives in {STATE_LABELS.get(_s(d, 'buyer_state'), 'another state')}. "
                     "The Seller shall give a no-objection certificate in Form 28 or help obtain it. The Buyer shall apply for a new registration mark "
                     "in the Buyer's state within 12 months of taking the Vehicle there (Section 47).")
    bearer = _s(d, "transfer_cost_bearer", "Buyer")
    items.append(f"Fees, taxes and endorsement charges for the transfer shall be borne by the {bearer}.")
    items.append("Until the transfer is recorded, the Buyer alone is responsible for the Vehicle's use and shall indemnify the Seller "
                 "against claims, challans and penalties arising from its use after delivery.")
    items.append("If the Buyer does not apply within 30 days, the Buyer authorises the Seller to report the transfer to the registering "
                 "authority. This does not reduce the Buyer's duties under this Clause.")
    sub(items)

    # 8 tax and challans
    clause("Tax, challans and dues")
    L.append(f"Tax under {S['tax_act']}, challans, fines and other dues for the period up to delivery are the Seller's. "
             "Those for the period after delivery are the Buyer's, whether or not the Vehicle has yet been transferred in the Buyer's name.")

    # 9 hypothecation
    clause("Hypothecation and finance")
    if hypo:
        L.append(f"The Vehicle is hypothecated to {_s(d, 'financier') or 'a financier'}. The Seller shall, at its cost, close the loan, "
                 "obtain Form 35 and the no-objection certificate, and have the hypothecation entry removed from the registration "
                 "certificate before delivery. If the Seller cannot do so before delivery, the Buyer may withhold from the Balance "
                 "the amount outstanding on the loan and pay it directly to the financier.")
    else:
        L.append("The Seller states that no hypothecation subsists on the Vehicle.")
    if _s(d, "finance_by_buyer") == "yes":
        L.append("The hypothecation in favour of the Buyer's financier shall be recorded in the manner the Motor Vehicles Act, 1988 requires, and the Seller shall cooperate.")

    # 10 insurance
    clause("Insurance")
    L.append("The Buyer shall apply to the insurer within 14 days of delivery to record the transfer of the policy (Section 157, Motor Vehicles Act, 1988), "
             "or buy a new policy, and shall not drive the Vehicle without valid insurance.")

    # 11 default
    clause("Default and termination")
    sub([f"If the Buyer does not pay the Balance by the due date, the Buyer shall pay interest on it at {rate}% a year from the due date until payment.",
         f"If the delay passes 15 days, the Seller may end this Agreement by written notice, keep from the advance, as reasonable compensation (Section 74, Indian Contract Act, 1872), "
         f"the lower of {pct}% of the Sale Price and the advance actually paid, and refund the rest. The Seller may then resell the Vehicle and claim any further loss "
         "(Section 54, Sale of Goods Act, 1930). If the Vehicle has been delivered, the Seller may instead sue for the price (Section 55 of that Act).",
         "If the Seller does not deliver within 15 days of the delivery date, for a reason not caused by the Buyer, the Buyer may end this Agreement by written notice and "
         "the Seller shall refund everything it has received within 7 days of that notice."])

    # 12 warranty
    clause("Limited warranty")
    wd, wk = int(num(d.get("warranty_days"))), int(num(d.get("warranty_km")))
    if wd > 0 or wk > 0:
        span = " or ".join(x for x in (f"{wd} days" if wd else "", f"{wk:,} km" if wk else "") if x)
        L.append(f"The Seller warrants the {_s(d, 'warranty_scope') or 'engine and gearbox'} for {span} from delivery, whichever ends first. "
                 "The Buyer shall report a defect in writing and bring the Vehicle to the Seller's workshop. The Seller shall repair or replace the faulty part at its cost. "
                 "The warranty does not cover accident damage, misuse, tyres, batteries, or parts that wear in normal use.")
    else:
        L.append("The Seller gives no warranty beyond the assurances in Clause 5, to the extent the law allows.")

    # 13 indemnity, 14 notices
    clause("Indemnity")
    L.append("Each party shall compensate the other for loss caused by a breach of its assurances or duties under this Agreement.")
    clause("Notices")
    L.append("Formal notices shall be in writing and sent by registered post or speed post to the addresses in this Agreement. "
             "Messages by email or chat apps are for convenience and are not formal notice.")

    # 15 law
    clause("Governing law and disputes")
    L.append("This Agreement is governed by the laws of India. Nothing in it limits any right the Seller has under the Negotiable Instruments Act, 1881 for a dishonoured cheque.")
    if str(d.get("arbitration")) in ("1", "true", "True", "yes"):
        L.append(f"Disputes shall be referred to a sole arbitrator chosen by agreement of both parties. If they cannot agree within 30 days of a written request, "
                 f"the arbitrator shall be appointed by {S['hc']} under Section 11 of the Arbitration and Conciliation Act, 1996. The seat and venue are {jur}, "
                 f"the language is English, and the courts at {jur} have jurisdiction over interim relief and enforcement.")
    else:
        L.append(f"The courts at {jur}, {S['name']} have exclusive jurisdiction over any dispute arising from this Agreement.")

    # 16 stamp, 17 language, 18 general
    clause("Stamp duty")
    sv = _s(d, "stamp_value")
    L.append(f"This Agreement is executed on stamp paper or e-stamp of value {inr(sv) if sv else 'Rs. ______'} bought in {S['name']}, "
             f"as required under {S['stamp_act']}. The duty is borne equally by both parties.")
    clause("Language")
    L.append(f"This Agreement is written in English. Its contents were read out and explained in {S['lang']} to the Buyer, "
             "who confirms that the Buyer understood them and signs of the Buyer's own free will.")
    clause("General")
    L.append("This Agreement is the whole agreement between the parties. A change is valid only if written and signed by both. "
             "If any part is held invalid, the rest continues. Time is of the essence for payment and for transfer of ownership.")
    extra = [x.strip() for x in _s(d, "extra_clauses").splitlines() if x.strip()]
    if extra:
        clause("Additional terms")
        sub(extra)

    L += ["", f"Signed at {place} on {agr_date}.", ""]
    seller_sign = p.get("signatory") or "Authorised signatory"
    L += [f"|~**For {_cell(p['dealership_name'])} (Seller)**|~**Buyer**",
          "|~|~", "|~|~",
          f"|~{_cell(seller_sign)}|~{_cell(_s(d, 'buyer_name'))}", "|~|~",
          f"|~**Witness 1:** {_cell(_s(d, 'witness1')) or '____________________'}|~**Witness 2:** {_cell(_s(d, 'witness2')) or '____________________'}",
          "|~Address and signature:|~Address and signature:"]

    # annexure
    L += ["~~", "## ANNEXURE A", "## DELIVERY AND DOCUMENT HANDOVER RECEIPT", "",
          f"Vehicle: {vehicle_long(d)}", "",
          "|**Item**|**Handed over (Yes / No)**|**Remarks**"]
    items = ["Original registration certificate", "Form 29 (two sets) and Form 30", "Insurance policy",
             "Pollution-under-control certificate", "Tax payment proof", "Service history", "Keys (number: ____)",
             "Spare wheel and tool kit"]
    if hypo:
        items.insert(2, "Form 35 and financier's no-objection certificate")
    if cross:
        items.insert(2, "Form 28 no-objection certificate")
    for it in items:
        L.append(f"|{it}||")
    L += ["", "Odometer at delivery: ________ km        Fuel level: ________", "",
          "The Buyer confirms receiving the Vehicle and the items ticked above on ________________ at ____________, in the condition inspected.", "",
          f"|~**Buyer:** {_cell(_s(d, 'buyer_name'))}|~**For {_cell(p['dealership_name'])}:** {_cell(seller_sign)}",
          "|~Signature:|~Signature:"]
    return "\n".join(L)


# ---------------------------------------------------------------- notices

def _f(extra, k, default=""):
    v = extra.get(k)
    if v in (None, ""):
        return default
    return _one_line(v, "; " if k == "challan_details" else " ")


def _field(key, label, kind="text", default="", req=False, hint=""):
    return {"key": key, "label": label, "type": kind, "default": default, "required": req, "hint": hint}


NOTICE_TYPES = {
    "payment_default": {
        "label": "Balance payment not made",
        "blurb": "The buyer has not paid the balance by the due date.",
        "fields": [_field("due_amount", "Amount still due (Rs.)", "number", req=True),
                   _field("due_date", "It was due on", "date", req=True),
                   _field("demand_days", "Days to pay after receiving the notice", "number", "15", True),
                   _field("interest_rate", "Interest (% a year)", "number", "12"),
                   _field("notice_cost", "Notice cost to claim (Rs., optional)", "number", "0"),
                   _field("extra_facts", "Anything else the advocate should know", "textarea")],
    },
    "cheque_dishonour": {
        "label": "Cheque bounced (Section 138)",
        "blurb": "The buyer's cheque came back unpaid. Strict deadlines apply.",
        "fields": [_field("cheque_no", "Cheque number", req=True),
                   _field("cheque_date", "Cheque date", "date", req=True),
                   _field("cheque_amount", "Cheque amount (Rs.)", "number", req=True),
                   _field("drawee_bank", "Buyer's bank and branch", req=True),
                   _field("dealer_bank", "Your bank where you deposited it", req=True),
                   _field("return_memo_date", "Date on the bank's return memo", "date", req=True),
                   _field("return_reason", "Reason on the return memo", "text", "Funds insufficient", True),
                   _field("intimation_date", "Date the bank told you it bounced", "date", "",
                          True, "The 30-day limit to send the notice runs from this date."),
                   _field("extra_facts", "Anything else the advocate should know", "textarea")],
    },
    "transfer_default": {
        "label": "RC not transferred",
        "blurb": "The buyer has not moved the registration into their name.",
        "fields": [_field("delivered_on", "Car was delivered on", "date", req=True),
                   _field("demand_days", "Days to complete transfer", "number", "15", True),
                   _field("challans", "Challans or notices already received (optional)", "textarea"),
                   _field("extra_facts", "Anything else the advocate should know", "textarea")],
    },
    "cancellation_forfeiture": {
        "label": "Buyer backed out",
        "blurb": "Final chance to complete the purchase, then cancel and keep part of the advance.",
        "fields": [_field("backed_out_on", "Date the buyer said they would not proceed", "date"),
                   _field("demand_days", "Days to pay the balance and take delivery", "number", "7", True),
                   _field("forfeit_amount", "Amount of advance you will keep (Rs.)", "number", req=True,
                          hint="Keep it to a reasonable share of the price. Courts do not enforce excessive forfeiture."),
                   _field("refund_days", "Days to refund the rest", "number", "15", True),
                   _field("extra_facts", "Anything else the advocate should know", "textarea")],
    },
    "challan_liability": {
        "label": "Challans after delivery",
        "blurb": "Traffic fines on the car after it left your lot, still showing against you.",
        "fields": [_field("delivered_on", "Car was delivered on", "date", req=True),
                   _field("challan_details", "Challan numbers, dates and amounts", "textarea", req=True),
                   _field("challan_total", "Total (Rs.)", "number", req=True),
                   _field("demand_days", "Days to pay and send proof", "number", "10", True),
                   _field("extra_facts", "Anything else the advocate should know", "textarea")],
    },
    "reply_to_complaint": {
        "label": "Reply to a buyer's complaint",
        "blurb": "The buyer says the car is defective. Respond on record.",
        "fields": [_field("complaint_date", "Date of the buyer's complaint or notice", "date", req=True),
                   _field("complaint_summary", "What the buyer alleges", "textarea", req=True),
                   _field("delivered_on", "Car was delivered on", "date"),
                   _field("km_run", "Kilometres run since delivery (if known)", "number"),
                   _field("extra_facts", "Anything else the advocate should know", "textarea")],
    },
}


def _header(p, d, st, mode_line):
    S = STATES[st]
    return ["# LEGAL NOTICE", "", f">> Date: {dmy(today_ist().isoformat())}", f">> {mode_line}", "",
            "To,", f"**{_s(d, 'buyer_name')}**",
            f"{_s(d, 'buyer_rel', 'S/o')} {_s(d, 'buyer_rel_name')}".strip(),
            ", ".join(x for x in (_s(d, "buyer_address"), _s(d, "buyer_city")) if x),
            (f"Mobile: {_s(d, 'buyer_phone')}" if _s(d, "buyer_phone") else ""),
            (f"Email: {_s(d, 'buyer_email')}" if _s(d, "buyer_email") else ""), ""]


def _opening(p, st):
    S = STATES[st]
    return ("Sir/Madam,", "",
            f"Under instructions from and on behalf of my client, **{p['dealership_name']}**, a dealer in pre-owned motor vehicles "
            f"with its place of business at {_s(p, 'address')}, {p['city']}, {S['name']} (\"my client\"), I give you this notice:")


def build_notice(ntype, p, d, extra, agreement_date=""):
    st = deal_state(p, d)
    S = STATES[st]
    jur = _s(d, "jurisdiction_city") or _s(d, "place") or p.get("city") or S["default_city"]
    veh = vehicle_long(d)
    price, adv = num(d.get("price")), num(d.get("advance"))
    agr = dmy(d.get("agreement_date")) or dmy(agreement_date) or "the date recorded in the agreement"
    days = int(num(_f(extra, "demand_days"), 15)) or 15
    cost = int(num(_f(extra, "notice_cost"), 0))
    L = []
    n = [0]

    def para(t):
        n[0] += 1
        L.append(f"{n[0]}. {t}")

    mode = "BY REGISTERED POST A.D. AND SPEED POST" + (" AND EMAIL" if _s(d, "buyer_email") else "")
    if ntype == "cheque_dishonour":
        title = "Notice under Section 138 of the Negotiable Instruments Act, 1881"
    else:
        title = {"payment_default": "Demand for payment of the balance sale price",
                 "transfer_default": "Demand to complete transfer of ownership of the vehicle",
                 "cancellation_forfeiture": "Final notice before cancellation of the sale",
                 "challan_liability": "Demand to pay traffic challans incurred after delivery",
                 "reply_to_complaint": "Reply to your complaint about the vehicle"}[ntype]
    L += _header(p, d, st, mode)
    L += [f"### Subject: {title}", "", *_opening(p, st)]
    L.append("")

    sale = (f"By an agreement dated {agr}, my client agreed to sell, and you agreed to buy, the vehicle {veh} (the \"vehicle\") "
            f"for {inr_words(price)}.")

    if ntype == "payment_default":
        due = num(_f(extra, "due_amount"), 0) or max(0, price - adv)
        rate = _f(extra, "interest_rate", "12")
        delivered = d.get("_stage", 0) >= 2
        para(sale)
        para((f"You paid an advance of {inr(adv)}. The balance of " if adv > 0 else "The sale price of ")
             + f"{inr_words(due)} was payable by {dmy(_f(extra, 'due_date')) or 'the agreed date'}.")
        para("Despite repeated requests, you have not paid the balance.")
        para("The vehicle was delivered to you and has been in your use since." if delivered else
             "My client continues to hold the vehicle ready for delivery against payment of the balance.")
        if _f(extra, "extra_facts"):
            para(_f(extra, "extra_facts"))
        para(f"I call upon you to pay my client {inr(due)} with interest at {rate}% a year from the due date until payment, within {days} days of receiving this notice."
             + (f" You shall also pay {inr(cost)} as the cost of this notice." if cost else ""))
        if delivered:
            para("If you do not, my client will file a civil suit against you for the price under Section 55 of the Sale of Goods Act, 1930, with interest and costs, at your risk as to cost and consequences.")
        else:
            para("If you do not, my client will treat this as notice of its intention to resell the vehicle under Section 54 of the Sale of Goods Act, 1930, "
                 "and will recover from you any loss on resale, with interest and costs.")

    elif ntype == "cheque_dishonour":
        amt = num(_f(extra, "cheque_amount"))
        para(sale)
        para(f"In part payment of that price and to discharge a legally enforceable debt, you issued to my client cheque no. {_f(extra, 'cheque_no')} dated "
             f"{dmy(_f(extra, 'cheque_date'))} for {inr_words(amt)} drawn on {_f(extra, 'drawee_bank')}.")
        para(f"My client presented the cheque through its bank, {_f(extra, 'dealer_bank')}. It was returned unpaid on {dmy(_f(extra, 'return_memo_date'))} "
             f"with the remark \"{_f(extra, 'return_reason')}\". My client received the bank's intimation on {dmy(_f(extra, 'intimation_date'))}.")
        para("The dishonour of the cheque is an offence under Section 138 of the Negotiable Instruments Act, 1881.")
        if _f(extra, "extra_facts"):
            para(_f(extra, "extra_facts"))
        para(f"I call upon you to pay my client {inr_words(amt)} within 15 days of receiving this notice. If you do not, my client will file a complaint against you "
             f"under Sections 138 and 142 of the Act, which carries imprisonment up to two years, a fine up to twice the cheque amount, or both, and will also recover the amount "
             f"in a civil suit with interest and costs.")

    elif ntype == "transfer_default":
        para(sale)
        para(f"The vehicle was delivered to you on {dmy(_f(extra, 'delivered_on'))}. The agreement required you to apply to the {S['authority_short']} within 30 days "
             f"of delivery for transfer of ownership into your name, as Section 50 of the Motor Vehicles Act, 1988 also requires. Forms 29 and 30 were signed and handed over for this purpose.")
        para("You have not completed the transfer. The vehicle therefore still stands in the name of its previous registered owner, which exposes my client to claims, challans and penalties for a vehicle it no longer holds.")
        if _f(extra, "challans"):
            para(f"The following have already been recorded against the vehicle: {_f(extra, 'challans')}")
        if _f(extra, "extra_facts"):
            para(_f(extra, "extra_facts"))
        para(f"I call upon you to apply for the transfer within {days} days of receiving this notice and to send my client the RTO acknowledgment, and to pay any challan that arose after delivery.")
        para("If you do not, my client will report the transfer to the registering authority, call on you to indemnify it under the agreement, and take legal action for the loss, with costs.")

    elif ntype == "cancellation_forfeiture":
        forfeit = num(_f(extra, "forfeit_amount"))
        refund = max(0.0, adv - forfeit)
        rdays = int(num(_f(extra, "refund_days"), 15)) or 15
        para(sale)
        para((f"You paid an advance of {inr(adv)}. You have not paid the balance of {inr(max(0, price - adv))}" if adv > 0
              else f"You have not paid the sale price of {inr(price)}") + " or taken delivery"
             + (f", and told my client on {dmy(_f(extra, 'backed_out_on'))} that you will not proceed." if _f(extra, "backed_out_on") else "."))
        if _f(extra, "extra_facts"):
            para(_f(extra, "extra_facts"))
        para(f"This is my client's last opportunity to complete the purchase. Pay the balance and take delivery within {days} days of receiving this notice. "
             f"This is also notice of my client's intention to resell the vehicle under Section 54 of the Sale of Goods Act, 1930 if you do not.")
        if adv > 0 and forfeit > 0:
            para(f"If you do not, the agreement will stand cancelled. My client will keep {inr(forfeit)} of the advance as reasonable compensation for its loss "
                 f"(Section 74, Indian Contract Act, 1872), refund {inr(refund)} within {rdays} days of receiving your bank details, and claim any further loss on resale.")
        else:
            para("If you do not, the agreement will stand cancelled. My client will claim reasonable compensation for its loss "
                 "(Section 74, Indian Contract Act, 1872), including any loss on resale of the vehicle.")

    elif ntype == "challan_liability":
        para(sale)
        para(f"The vehicle was delivered to you on {dmy(_f(extra, 'delivered_on'))} and has been in your possession and control since. "
             "The agreement makes you responsible for all challans, fines and dues arising after delivery.")
        para(f"The following challans have been recorded against the vehicle for the period after delivery: {_f(extra, 'challan_details')}. "
             f"The total is {inr(_f(extra, 'challan_total'))}.")
        if _f(extra, "extra_facts"):
            para(_f(extra, "extra_facts"))
        para(f"I call upon you to pay these challans and send my client the receipts within {days} days of receiving this notice, and to complete the transfer of the vehicle into your name.")
        para("If you do not, my client will pay them to protect its record, recover the amount from you with interest, and claim damages and costs.")

    elif ntype == "reply_to_complaint":
        para(f"I refer to your complaint dated {dmy(_f(extra, 'complaint_date'))} about the vehicle {veh}, which you bought from my client for {inr_words(price)} under an agreement dated {agr}. "
             f"You allege: {_f(extra, 'complaint_summary')}")
        para("My client denies your allegations. You inspected and test-drove the vehicle before buying. It was sold in the condition you inspected, "
             f"with its known defects disclosed ({_s(d, 'known_defects') or 'none other than normal wear and tear'}), and its odometer reading of {_s(d, 'odometer') or 'the recorded figure'} km was recorded.")
        km = _f(extra, "km_run")
        if _f(extra, "delivered_on") or km:
            para("The vehicle was delivered to you" + (f" on {dmy(_f(extra, 'delivered_on'))}" if _f(extra, "delivered_on") else "")
                 + (f" and has been driven about {km} km since." if km else "."))
        wd, wk = int(num(d.get("warranty_days"))), int(num(d.get("warranty_km")))
        if wd or wk:
            para("The agreement gives a limited warranty. If the fault you describe falls within it, my client will repair it at its workshop at no cost to you.")
        if _f(extra, "extra_facts"):
            para(_f(extra, "extra_facts"))
        para("I call upon you to bring the vehicle for a joint inspection on a date within 10 days of receiving this notice, and to send an independent workshop report and invoices for any repair you say you have paid for.")
        para(f"My client reserves all its rights and will defend any claim before any forum, including the District Consumer Disputes Redressal Commission at {jur}. "
             "A claim that proves false will be met with a claim for costs.")

    L += ["", f"Please treat this as notice under the agreement and the law. My client reserves all its other rights. A copy of this notice is kept in my office for record."]
    return title, "\n".join(x for x in L if x is not None)


def s138_deadline(extra):
    """Days left in the 30-day window to send a Section 138 notice, or None."""
    t = parse_date((extra or {}).get("intimation_date"))
    if not t:
        return None
    return 30 - (today_ist() - t).days


# ============================================================================
# 4. DEAL AND NOTICE VALIDATION
# ============================================================================

def _txt(v, n=300):
    return str(v if v is not None else "").strip()[:n]


def _safe(s):
    return re.sub(r"[^\w\-]+", "_", s or "").strip("_") or "document"


_DATE_RE = re.compile(r"\d{4}-\d{2}-\d{2}")


def valid_date(v):
    """True for a real calendar date written YYYY-MM-DD (what the date pickers send)."""
    return bool(_DATE_RE.fullmatch(v or "")) and parse_date(v) is not None


PROFILE_FIELDS = ("dealership_name", "signatory", "address", "city", "state", "gstin", "trade_cert", "phone")

DEAL_KEYS = [
    "buyer_name", "buyer_rel", "buyer_rel_name", "buyer_age", "buyer_address", "buyer_city", "buyer_state",
    "buyer_phone", "buyer_email", "buyer_id_type", "buyer_id_no",
    "make", "model", "variant", "year", "reg_no", "reg_state", "chassis", "engine_no", "colour", "fuel",
    "odometer", "owners", "rc_holder", "hypothecated", "financier", "insurance_policy", "insurance_valid",
    "known_defects", "accident_history",
    "price", "advance", "advance_date", "advance_mode", "balance_due_date", "balance_mode",
    "finance_by_buyer", "buyer_financier", "delivery_date", "warranty_days", "warranty_km", "warranty_scope",
    "transfer_cost_bearer", "interest_rate", "forfeit_pct",
    "deal_state", "place", "jurisdiction_city", "arbitration", "agreement_date", "stamp_value",
    "extra_clauses", "witness1", "witness2",
]
NUMERIC = {"buyer_age": "Buyer age", "year": "Year", "odometer": "Odometer", "price": "Sale price", "advance": "Advance",
           "warranty_days": "Warranty days", "warranty_km": "Warranty km", "interest_rate": "Interest rate",
           "forfeit_pct": "Forfeiture percentage", "stamp_value": "Stamp paper value"}
LONG_TEXT = {"buyer_address", "known_defects", "extra_clauses"}
DATE_KEYS = {"advance_date": "Advance date", "balance_due_date": "Balance due date", "delivery_date": "Delivery date",
             "insurance_valid": "Insurance validity date", "agreement_date": "Agreement date"}


def clean_profile(j):
    p, errs = {}, []
    for k in PROFILE_FIELDS:
        p[k] = _txt(j.get(k), 30 if k == "phone" else 500)
    for k, label in (("dealership_name", "dealership name"), ("address", "dealership address"), ("city", "city")):
        if not p[k]:
            errs.append(f"Enter your {label}.")
    if p["state"] not in STATES:
        errs.append("Choose Tamil Nadu or Kerala for the dealership.")
    return p, errs


def clean_deal(j):
    d, errs = {}, []
    for k in DEAL_KEYS:
        v = _txt(j.get(k), 2000 if k in LONG_TEXT else 200)
        if k in NUMERIC and v:
            v = re.sub(r"[,\s]", "", v)
            if not re.fullmatch(r"\d+(\.\d+)?", v):
                errs.append(f"{NUMERIC[k]} must be a number.")
        d[k] = v
    for k in ("reg_no", "chassis", "engine_no"):
        d[k] = re.sub(r"\s+", " ", d[k]).upper()
    # Aadhaar numbers are never kept in full: only the last four digits are used in the documents.
    digits = re.sub(r"\D", "", d["buyer_id_no"])
    id_type = d["buyer_id_type"].lower()
    if id_type.startswith("aadhaar") or (len(digits) == 12 and not re.search(r"[A-Za-z]", d["buyer_id_no"].replace("X", "").replace("x", ""))):
        if len(digits) >= 4:
            d["buyer_id_no"] = "XXXX XXXX " + digits[-4:]
        elif "X" not in d["buyer_id_no"].upper():
            d["buyer_id_no"] = ""
    for k, label in (("buyer_name", "buyer name"), ("buyer_address", "buyer address"), ("make", "make"),
                     ("model", "model"), ("reg_no", "registration number"), ("price", "sale price")):
        if not d[k]:
            errs.append(f"Enter the {label}.")
    for k, label in DATE_KEYS.items():
        if d[k] and not valid_date(d[k]):
            errs.append(f"{label} is not a valid date.")
    if d["year"] and re.fullmatch(r"\d+", d["year"]) and not 1950 <= int(d["year"]) <= today_ist().year + 1:
        errs.append("Year of manufacture looks wrong.")
    if d["buyer_age"] and re.fullmatch(r"\d+", d["buyer_age"]) and not 1 <= int(d["buyer_age"]) <= 120:
        errs.append("Buyer age looks wrong.")
    for k in ("forfeit_pct", "interest_rate"):
        if d[k] and re.fullmatch(r"\d+(\.\d+)?", d[k]) and float(d[k]) > 100:
            errs.append(f"{NUMERIC[k]} cannot be more than 100.")
    if not errs:
        if num(d["price"]) <= 0:
            errs.append("The sale price must be more than zero.")
        if num(d["advance"]) > num(d["price"]):
            errs.append("The advance cannot be more than the sale price.")
    if d["deal_state"] not in STATES:
        errs.append("Choose the state this deal is made in.")
    return d, errs


def notice_prefill(d):
    price, adv = num(d.get("price")), num(d.get("advance"))
    pct = num(d.get("forfeit_pct"), 10)
    return {"due_amount": int(max(0, price - adv)), "due_date": d.get("balance_due_date", ""),
            "delivered_on": d.get("delivery_date", ""), "forfeit_amount": int(min(adv, price * pct / 100)),
            "cheque_amount": int(max(0, price - adv))}


def clean_facts(ntype, raw, d):
    """Validate the details typed for one notice. Returns (facts, error_message)."""
    label = NOTICE_TYPES[ntype]["label"]
    if not isinstance(raw, dict):
        return None, f"{label}: the details were not understood."
    facts = {}
    for f in NOTICE_TYPES[ntype]["fields"]:
        v = _txt(raw.get(f["key"]), 3000 if f["type"] == "textarea" else 200) or f["default"]
        if f["type"] == "number" and v:
            v = re.sub(r"[,\s]", "", v)
            if not re.fullmatch(r"\d+(\.\d+)?", v):
                return None, f"{label}: {f['label']} must be a number."
            if f["key"].endswith("_days") and not 1 <= float(v) <= 365:
                return None, f"{label}: {f['label']} must be between 1 and 365."
        if f["type"] == "date" and v and not valid_date(v):
            return None, f"{label}: {f['label']} is not a valid date."
        if f["required"] and not v:
            return None, f"{label}: fill in \"{f['label']}\"."
        facts[f["key"]] = v
    if ntype == "cancellation_forfeiture" and num(facts.get("forfeit_amount")) > num(d.get("advance")):
        return None, f"{label}: you can't keep more than the advance you received ({inr(d.get('advance'))})."
    return facts, None


# ============================================================================
# 5. AI DEAL REVIEW AND CLARIFYING QUESTIONS
#    Rule-based checks always run. The AI adds a written review and extra questions when GROQ_API_KEY is set.
# ============================================================================

def deal_summary(p, d):
    """Plain-text facts for the AI. The buyer's name, address, phone, email and ID number are left out on purpose."""
    st = deal_state(p, d)
    S = STATES[st]
    price, adv = num(d.get("price")), num(d.get("advance"))
    rows = [
        ("Deal made in", S["name"]), ("Dealership city", p.get("city")),
        ("Vehicle", vehicle_label(d)), ("Registered in", STATE_LABELS.get(_s(d, "reg_state"), _s(d, "reg_state"))),
        ("Odometer (km)", _s(d, "odometer")), ("Owner serial", _s(d, "owners")), ("Fuel", _s(d, "fuel")),
        ("Loan on the car (hypothecated)", _s(d, "hypothecated") or "no"), ("Financier", _s(d, "financier")),
        ("Insurance valid until", dmy(d.get("insurance_valid"))),
        ("Accident history", _s(d, "accident_history") or "none disclosed"),
        ("Known defects", _s(d, "known_defects") or "none stated"),
        ("Sale price", inr(price)), ("Advance", f"{inr(adv)} by {_s(d, 'advance_mode') or 'unspecified mode'}"),
        ("Balance", f"{inr(balance_due(d))} by {_s(d, 'balance_mode') or 'unspecified mode'}, due {dmy(d.get('balance_due_date')) or 'not set'}"),
        ("Buyer is financing the balance", _s(d, "finance_by_buyer") or "no"),
        ("Delivery date", dmy(d.get("delivery_date")) or "not set"),
        ("Buyer lives in", STATE_LABELS.get(_s(d, "buyer_state"), _s(d, "buyer_state"))), ("Buyer age", _s(d, "buyer_age")),
        ("Warranty", f"{_s(d, 'warranty_days') or 0} days / {_s(d, 'warranty_km') or 0} km"),
        ("Forfeiture on default (% of price)", _s(d, "forfeit_pct") or "10"),
        ("Interest on late balance (% a year)", _s(d, "interest_rate") or "12"),
        ("Transfer cost borne by", _s(d, "transfer_cost_bearer") or "Buyer"),
        ("Stamp paper value", _s(d, "stamp_value") or "not set"), ("Witnesses named", "yes" if _s(d, "witness1") else "no"),
        ("Arbitration clause", "yes" if _s(d, "arbitration") else "no"),
    ]
    return "\n".join(f"- {k}: {v}" for k, v in rows if v not in (None, ""))


def _parse_json(text):
    text = re.sub(r"^```\w*\s*|\s*```$", "", (text or "").strip())
    starts = [i for i in (text.find("{"), text.find("[")) if i >= 0]
    if not starts:
        raise ValueError("no JSON in reply")
    return json.JSONDecoder().raw_decode(text[min(starts):])[0]


def _clip(v, n=300):
    return " ".join(str(v or "").split())[:n]


def _clip_list(v, count, size=300):
    return [x for x in (_clip(i, size) for i in (v if isinstance(v, list) else [])) if x][:count]


RISK_ORDER = {"low": 0, "medium": 1, "high": 2}

INSIGHT_SYSTEM = (
    "You advise a used-car dealership in Tamil Nadu or Kerala, India, in plain English. You are given the facts of one sale "
    "and the rule-based checks already run. Give the dealer a practical review of the deal. Use only the facts given. Never invent "
    "facts, figures, section numbers or case law. You are not a lawyer: where the law is involved, say an advocate should confirm. "
    'Reply with JSON only, no markdown, in exactly this shape: {"verdict": "two sentences on how sound the deal looks", '
    '"risk_level": "low or medium or high", "strengths": ["..."], "risks": [{"issue": "...", "action": "..."}], "next_steps": ["..."]}. '
    "At most 5 strengths, 6 risks and 6 next steps. Keep each item under 40 words."
)


def rule_based_review(p, d, comp):
    """Deterministic review: numbers, risk level, and the risks/next steps that follow from the checks."""
    st = deal_state(p, d)
    S = STATES[st]
    price, adv, bal = num(d.get("price")), num(d.get("advance")), balance_due(d)
    pct = num(d.get("forfeit_pct"), 10)
    yr = int(num(d.get("year"))) if num(d.get("year")) else 0
    adv_pct = (adv / price * 100) if price else 0
    metrics = [
        {"label": "Sale price", "value": inr(price)},
        {"label": "Advance received", "value": f"{inr(adv)} ({adv_pct:.0f}% of price)"},
        {"label": "Balance due", "value": inr(bal)},
        {"label": "Most you can keep if the buyer defaults", "value": inr(min(adv, price * pct / 100))},
        {"label": "Vehicle age", "value": f"{today_ist().year - yr} years" if yr else "Year not given"},
        {"label": "Deal state", "value": S["name"]},
    ]
    alerts = sum(1 for c in comp if c["level"] == "alert")
    warns = sum(1 for c in comp if c["level"] == "warn")
    risks = [{"issue": c["text"], "action": ""} for c in comp if c["level"] in ("alert", "warn")]
    if price and adv <= 0:
        risks.append({"issue": "No advance has been taken, so there is nothing to keep if the buyer walks away.",
                      "action": "Take a token advance and record it in the agreement."})
        warns += 1
    elif price and adv_pct < 10:
        risks.append({"issue": f"The advance is only {adv_pct:.0f}% of the price, so a forfeiture claim on default would be small.",
                      "action": "Ask for a larger advance before holding the car for the buyer."})
        warns += 1
    level = "high" if alerts >= 2 else ("medium" if alerts == 1 or warns >= 2 else "low")
    verdict = {"low": "No serious problems found in the details you entered. Check the points below and the documents before signing.",
               "medium": "The deal is workable, but a few points need attention before you hand over the car.",
               "high": "Several points need fixing before this deal is safe to close. Deal with the red items first."}[level]
    steps = ["Have the advocate you work with read the agreement before first use.",
             "Collect the buyer's identity and address proof and keep copies.",
             "Hand over Forms 29 and 30 only against payment, and keep a signed receipt."]
    return {"verdict": verdict, "risk_level": level, "metrics": metrics, "compliance": comp,
            "strengths": [], "risks": risks[:6], "next_steps": steps}


def ai_review(p, d, comp):
    """Returns (ai_fields or None, note)."""
    if not ai_available():
        return None, "The AI is switched off on this server (GROQ_API_KEY is not set), so you are seeing the rule-based review only."
    prompt = ("DEAL FACTS\n" + deal_summary(p, d) + "\n\nRULE-BASED CHECKS ALREADY RUN\n" +
              "\n".join(f"- [{c['level']}] {c['text']}" for c in comp))
    try:
        data = _parse_json(ai_generate(prompt, system=INSIGHT_SYSTEM, temperature=0.2, max_tokens=1800))
        if not isinstance(data, dict):
            raise ValueError("unexpected shape")
    except Exception as e:                                   # the review still works without the AI
        print(f"[AI] review failed: {e}")
        return None, "The AI could not finish the review just now, so you are seeing the rule-based review only."
    risks = []
    for r in (data.get("risks") if isinstance(data.get("risks"), list) else [])[:6]:
        if isinstance(r, dict) and _clip(r.get("issue")):
            risks.append({"issue": _clip(r.get("issue")), "action": _clip(r.get("action"))})
    return {"verdict": _clip(data.get("verdict"), 500), "risk_level": str(data.get("risk_level", "")).lower(),
            "strengths": _clip_list(data.get("strengths"), 5), "risks": risks,
            "next_steps": _clip_list(data.get("next_steps"), 6)}, ""


# --- clarifying questions ------------------------------------------------------------------
# The AI may only ask about these fields, so an answer can only ever fill a known slot of the deal.

YESNO_ARB = [{"value": "", "label": "No arbitration clause (courts decide)"}, {"value": "yes", "label": "Add an arbitration clause"}]
Q_FIELDS = {
    "stamp_value": ("Value of the stamp paper or e-stamp (Rs.)", "number", None),
    "delivery_date": ("Date you will deliver the car", "date", None),
    "balance_due_date": ("Date the balance is due", "date", None),
    "insurance_valid": ("Insurance valid until", "date", None),
    "chassis": ("Chassis number (from the RC)", "text", None),
    "engine_no": ("Engine number (from the RC)", "text", None),
    "financier": ("Which bank or financier holds the loan on the car?", "text", None),
    "buyer_financier": ("Which lender is financing the buyer?", "text", None),
    "known_defects": ("Known defects to disclose to the buyer (leave blank if none)", "textarea", None),
    "witness1": ("Name of witness 1", "text", None),
    "witness2": ("Name of witness 2", "text", None),
    "warranty_scope": ("What does your warranty cover?", "text", None),
    "transfer_cost_bearer": ("Who pays the RC transfer fees?", "select", [{"value": "Buyer", "label": "Buyer"}, {"value": "Seller", "label": "Seller (dealership)"}]),
    "arbitration": ("Should disputes go to arbitration?", "select", YESNO_ARB),
    "extra_clauses": ("Any special term to add (one per line)", "textarea", None),
}
NOTE = lambda k: {"financier": "The agreement names the lender, and Form 35 plus the lender's NOC are needed before handover.",
                  "buyer_financier": "The agreement makes the buyer answerable for the lender's payment dates.",
                  "balance_due_date": "Interest and default rights run from this date.",
                  "delivery_date": "Warranty, risk and transfer deadlines are counted from delivery.",
                  "stamp_value": "Your advocate confirms the value. Leave blank to print a blank to fill in.",
                  "insurance_valid": "We warn you if the policy ends before delivery.",
                  "chassis": "Chassis and engine numbers identify the car and stop disputes about which car was sold.",
                  "engine_no": "Chassis and engine numbers identify the car and stop disputes about which car was sold.",
                  "known_defects": "A defect that is written down is a defect the buyer accepted.",
                  "witness1": "Named witnesses make the agreement easier to prove.",
                  "witness2": "Named witnesses make the agreement easier to prove."}.get(k, "")


def gap_questions(d):
    """Rule-based questions about things the agreement needs and the deal does not yet say."""
    gaps = []

    def need(k, cond=True):
        if cond and not _s(d, k):
            gaps.append(k)
    need("financier", _s(d, "hypothecated") == "yes")
    need("buyer_financier", _s(d, "finance_by_buyer") == "yes")
    need("balance_due_date", balance_due(d) > 0)
    need("delivery_date")
    need("stamp_value")
    need("insurance_valid")
    need("chassis")
    need("engine_no")
    need("known_defects", _s(d, "accident_history") in ("minor", "major"))
    need("witness1")
    return gaps[:6]


QUESTIONS_SYSTEM = (
    "You help prepare a used-car sale agreement for a dealership in Tamil Nadu or Kerala, India. Given the deal facts, ask at most 2 "
    "short clarifying questions whose answers would make the agreement more accurate or safer. Ask only about things the facts do not already answer. "
    "Each question must map to exactly one of these field ids: " + ", ".join(Q_FIELDS) + ". "
    'Reply with JSON only: [{"field": "<field id>", "question": "<plain-English question>", "why": "<one short sentence>"}]. '
    "If nothing important is missing, reply with []."
)


def build_questions(p, d):
    out, seen = [], set()
    for k in gap_questions(d):
        label, kind, opts = Q_FIELDS[k]
        out.append({"field": k, "label": label, "type": kind, "options": opts, "why": NOTE(k), "ai": False})
        seen.add(k)
    if ai_available():
        try:
            data = _parse_json(ai_generate("DEAL FACTS\n" + deal_summary(p, d) + "\n\nQuestions already planned about: " + ", ".join(sorted(seen)),
                                           system=QUESTIONS_SYSTEM, temperature=0.2, max_tokens=500))
            for q in (data if isinstance(data, list) else [])[:3]:
                k = q.get("field") if isinstance(q, dict) else None
                if k in Q_FIELDS and k not in seen and (k == "extra_clauses" or not _s(d, k)) and len(out) < 8:
                    label, kind, opts = Q_FIELDS[k]
                    out.append({"field": k, "label": _clip(q.get("question"), 200) or label, "type": kind, "options": opts,
                                "why": _clip(q.get("why"), 200), "ai": True})
                    seen.add(k)
        except Exception as e:
            print(f"[AI] questions failed: {e}")
    return out


def suggest_notices(d):
    """Notices that are worth checking today, from the dates in the deal. The dealer decides."""
    today, out = today_ist(), []
    due, dlv = parse_date(d.get("balance_due_date")), parse_date(d.get("delivery_date"))
    if due and balance_due(d) > 0 and due < today:
        out.append({"key": "payment_default", "reason": f"The balance was due on {dmy(d.get('balance_due_date'))}."})
    if dlv and (today - dlv).days > 30:
        out.append({"key": "transfer_default", "reason": "More than 30 days have passed since delivery. Check the RC has moved to the buyer."})
    return out


def apply_answers(raw, answers):
    """Merge the dealer's answers into the deal. Only whitelisted fields, each checked by its type."""
    errs = []
    if not isinstance(answers, dict):
        return errs
    for k, v in answers.items():
        if k not in Q_FIELDS:
            continue
        label, kind, opts = Q_FIELDS[k]
        v = _txt(v, 2000)
        if not v:
            continue
        if kind == "date" and not valid_date(v):
            errs.append(f"\"{label}\" is not a valid date.")
            continue
        if kind == "number" and not re.fullmatch(r"[\d,\s]+(\.\d+)?", v):
            errs.append(f"\"{label}\" must be a number.")
            continue
        if kind == "select" and v not in [o["value"] for o in opts]:
            continue
        if k == "extra_clauses":
            raw[k] = (str(raw.get(k) or "").strip() + "\n" + v).strip()
        else:
            raw[k] = v
    return errs


# ============================================================================
# 6. WEB ROUTES
# ============================================================================

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
APP_NAME = "MFA DealPick"
DOCX_MIME = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
BANNER = "DRAFT FOR ADVOCATE REVIEW. NOT LEGAL ADVICE."
ADVOCATE_LINE = ">> [Advocate's signature and Bar Council enrolment number to be added after review.]"
DEMO_USER = (os.environ.get("DEALER_USERNAME") or "demo").strip()
DEMO_PASS = os.environ.get("DEALER_PASSWORD") or "demo1234"
if not os.environ.get("DEALER_PASSWORD"):
    print("[DealPick] DEALER_PASSWORD is not set. The login is demo / demo1234. Set your own before sharing the link.")

app = Flask(__name__)
_secret = os.environ.get("SECRET_KEY", "").strip()
if not _secret:
    _secret = secrets.token_hex(32)
    print("[DealPick] SECRET_KEY is not set. Using a temporary key; everyone is logged out on restart.")
app.config.update(
    SECRET_KEY=_secret,
    MAX_CONTENT_LENGTH=1024 * 1024,
    SESSION_COOKIE_HTTPONLY=True,
    SESSION_COOKIE_SAMESITE="Lax",
    SESSION_COOKIE_SECURE=bool(os.environ.get("RENDER")),
    PERMANENT_SESSION_LIFETIME=dt.timedelta(hours=12),
)
if os.environ.get("RENDER"):
    app.wsgi_app = ProxyFix(app.wsgi_app, x_for=1, x_proto=1, x_host=1)

try:
    with open(os.path.join(BASE_DIR, "index.html"), encoding="utf-8") as _fh:
        INDEX_HTML = _fh.read()
except OSError:
    INDEX_HTML = "<h1>index.html is missing</h1><p>Keep index.html in the same folder as app.py.</p>"


def fail(msg, code=400):
    return jsonify({"success": False, "message": msg}), code


def ok(**kw):
    kw["success"] = True
    return jsonify(kw)


@app.before_request
def _csrf_guard():
    # State-changing API calls must carry a header a cross-site form cannot set.
    if request.method in ("POST", "PUT", "PATCH", "DELETE") and request.path.startswith("/api/"):
        if request.headers.get("X-Requested-With") != "DealPick":
            return fail("Request blocked. Reload the page and try again.", 403)


@app.after_request
def _headers(resp):
    nonce = getattr(g, "nonce", "")
    resp.headers["X-Content-Type-Options"] = "nosniff"
    resp.headers["X-Frame-Options"] = "DENY"
    resp.headers["Referrer-Policy"] = "same-origin"
    resp.headers["Content-Security-Policy"] = (
        f"default-src 'self'; script-src 'nonce-{nonce}'; style-src 'nonce-{nonce}' https://fonts.googleapis.com; "
        "style-src-attr 'unsafe-inline'; font-src https://fonts.gstatic.com; img-src 'self' data:; "
        "connect-src 'self'; base-uri 'none'; form-action 'none'; frame-ancestors 'none'")
    if os.environ.get("RENDER"):
        resp.headers["Strict-Transport-Security"] = "max-age=31536000"
    if request.path.startswith("/api/"):
        resp.headers["Cache-Control"] = "no-store"
    return resp


@app.errorhandler(413)
def _too_large(_e):
    return fail("That request is too large.", 413)


@app.errorhandler(Exception)
def _on_error(e):
    if isinstance(e, HTTPException):
        if request.path.startswith("/api/"):
            return fail(e.description or e.name, e.code)
        return e
    traceback.print_exc()
    return fail("Something went wrong on the server. Please try again.", 500)


_FAILS = {}


def _throttled(key):
    t = time.time()
    if len(_FAILS) > 2000:
        for k in [k for k, v in _FAILS.items() if not v or t - v[-1] > 600]:
            _FAILS.pop(k, None)
    recent = [x for x in _FAILS.get(key, []) if t - x < 600]
    if recent:
        _FAILS[key] = recent
    else:
        _FAILS.pop(key, None)
    return len(recent) >= 6


def login_required(fn):
    def inner(*a, **kw):
        if not session.get("auth"):
            return fail("Please log in.", 401)
        return fn(*a, **kw)
    inner.__name__ = fn.__name__
    return inner


# --- login (a simple gate for the prototype; no accounts are stored) -----------------------

@app.post("/api/login")
def login():
    j = request.get_json(silent=True) or {}
    user, pw = _txt(j.get("username"), 100), str(j.get("password") or "")[:200]
    key = (request.remote_addr, user.lower())
    if _throttled(key):
        return fail("Too many tries. Wait 10 minutes and try again.", 429)
    good = hmac.compare_digest(user.lower().encode(), DEMO_USER.lower().encode()) & \
        hmac.compare_digest(pw.encode(), DEMO_PASS.encode())
    if not good:
        _FAILS.setdefault(key, []).append(time.time())
        return fail("Username or password is not right.", 401)
    session.clear()
    session["auth"] = True
    session.permanent = True
    return ok(ai=ai_available())


@app.post("/api/logout")
def logout():
    session.clear()
    return ok()


@app.get("/api/me")
def me():
    return ok(authed=bool(session.get("auth")), ai=ai_available())


# --- step 1 -> 2: review the deal ------------------------------------------------------------

def _read_deal(j):
    p, e1 = clean_profile(j.get("profile") if isinstance(j.get("profile"), dict) else {})
    d, e2 = clean_deal(j.get("deal") if isinstance(j.get("deal"), dict) else {})
    return p, d, e1 + e2


@app.post("/api/analyze")
@login_required
def analyze():
    p, d, errs = _read_deal(request.get_json(silent=True) or {})
    if errs:
        return fail(" ".join(errs))
    comp = compliance(p, d)
    review = rule_based_review(p, d, comp)
    ai, note = ai_review(p, d, comp)
    if ai:
        if ai["verdict"]:
            review["verdict"] = ai["verdict"]
        if ai["risk_level"] in RISK_ORDER and RISK_ORDER[ai["risk_level"]] > RISK_ORDER[review["risk_level"]]:
            review["risk_level"] = ai["risk_level"]                  # the AI can raise the level, never lower it
        review["strengths"] = ai["strengths"]
        if ai["risks"]:
            review["risks"] = ai["risks"]
        if ai["next_steps"]:
            review["next_steps"] = ai["next_steps"]
    return ok(review=review, ai=bool(ai), ai_note=note)


# --- step 2 -> 3: clarifying questions and the notices on offer ----------------------------------

@app.post("/api/questions")
@login_required
def questions():
    p, d, errs = _read_deal(request.get_json(silent=True) or {})
    if errs:
        return fail(" ".join(errs))
    prefill = notice_prefill(d)
    notices = [{"key": k, "label": v["label"], "blurb": v["blurb"], "fields": v["fields"]} for k, v in NOTICE_TYPES.items()]
    return ok(questions=build_questions(p, d), notices=notices, prefill=prefill, suggested=suggest_notices(d))


# --- step 3 -> 4: draft the documents --------------------------------------------------------------

@app.post("/api/generate")
@login_required
def generate():
    j = request.get_json(silent=True) or {}
    raw = dict(j.get("deal") if isinstance(j.get("deal"), dict) else {})
    errs = apply_answers(raw, j.get("answers"))
    p, d, more = _read_deal({"profile": j.get("profile"), "deal": raw})
    errs += more
    if errs:
        return fail(" ".join(errs))
    d["_stage"] = 2 if j.get("delivered") else 0                   # lets the payment notice say whether the car was delivered
    st, today = deal_state(p, d), today_ist()
    ref = f"DP-{st}-{today:%Y%m%d}-{_safe(d['reg_no'])}"
    docs = [{"key": "agreement", "title": "Sale agreement and delivery receipt", "body": build_agreement(p, d)}]
    wanted = j.get("notices") if isinstance(j.get("notices"), list) else []
    done = set()
    for item in wanted[:6]:
        if not isinstance(item, dict) or item.get("type") not in NOTICE_TYPES or item["type"] in done:
            continue
        ntype = item["type"]
        facts, err = clean_facts(ntype, item.get("facts") or {}, d)
        if err:
            return fail(err)
        title, body = build_notice(ntype, p, d, facts, d.get("agreement_date") or today.isoformat())
        docs.append({"key": ntype, "title": NOTICE_TYPES[ntype]["label"] + " (legal notice)", "body": body + "\n\n" + ADVOCATE_LINE})
        done.add(ntype)
    return ok(documents=docs, ref=ref)


@app.post("/api/download")
@login_required
def download():
    j = request.get_json(silent=True) or {}
    docs = j.get("documents")
    if not isinstance(docs, list) or not 1 <= len(docs) <= 8:
        return fail("There is nothing to download.")
    parts = []
    for x in docs:
        body = str(x.get("body") if isinstance(x, dict) else "").replace("\r\n", "\n").strip()
        if not body or len(body) > 60000:
            return fail("A document is empty or too long.")
        parts.append(body)
    ref = _safe(_txt(j.get("ref"), 80)) or "Deal"
    buf = build_docx("\n~~\n".join(parts), header_banner=BANNER, footer=f"{APP_NAME}   {ref}")
    return send_file(buf, as_attachment=True, download_name=f"Legal-Documents_{ref}.docx", mimetype=DOCX_MIME)


# --- pages -------------------------------------------------------------------------------------

@app.get("/healthz")
def healthz():
    return "ok", 200


@app.get("/")
def index():
    g.nonce = secrets.token_urlsafe(16)
    resp = Response(INDEX_HTML.replace("__NONCE__", g.nonce), mimetype="text/html")
    resp.headers["Cache-Control"] = "no-cache"
    return resp


if __name__ == "__main__":
    port = int(os.environ.get("PORT", 8081))
    print("\n" + "=" * 60)
    print(f"  {APP_NAME} (prototype, nothing is stored)")
    print(f"  Open: http://127.0.0.1:{port}")
    print(f"  Login: {DEMO_USER} / {'(your DEALER_PASSWORD)' if os.environ.get('DEALER_PASSWORD') else DEMO_PASS}")
    print(f"  AI review (GROQ_API_KEY): {'on' if ai_available() else 'off'}")
    print("=" * 60 + "\n")
    app.run(host="0.0.0.0", port=port, debug=False, threaded=True)
